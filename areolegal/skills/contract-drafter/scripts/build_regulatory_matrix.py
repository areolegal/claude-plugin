#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""קובץ ה-Excel של מצב 4, לחברה שעליה עובדים בהרצה הזאת.

    python3 build_regulatory_matrix.py --research "Areo-נתונים/research-<slug>.json" \\
        --out "<סוג ההסכם> - מטריצה רגולטורית <יום.חודש.שנה>.xlsx"

הגיליונות נקבעים לפי היקף העבודה שהמשתמש בחר (scope):
  research / both   "מטריצה רגולטורית" ו"סעיפי חובה"
  agreements / both "הסכמים מייצגים"

כל הגיליונות מימין לשמאל, שורת כותרת מוקפאת עם מסנן, Arial ועטיפת טקסט.
כל מקור ישירות לחיץ, וכל סעיף חובה מקושר פנימית לשורת המקור שלו. תאריכים
נכתבים כמחרוזת יום.חודש.שנה, כדי שהתצוגה האזורית לא תהפוך אותם.

קוד יציאה 1 כשכללי המקורות של research_sources.py נכשלים: מקור מדומיין אסור,
מקור בלי קישור, סדר שאינו לפי המדרג, או סעיף חובה על מקור לא מאומת.
"""
import argparse
import io
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from research_sources import (NORM_RANKS, RANK_STATUS, check, display_date,  # noqa: E402
                              normalise, scope_has)

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="12314F")
STATUS_FILL = {"מחייב": "E8EDF2", "מנחה": "F4F4F4", "בהליך": "FFF4E0"}
THIN = Side(style="thin", color="C9CFD6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)

LAW_SHEET = "מטריצה רגולטורית"
MAND_SHEET = "סעיפי חובה"
AGR_SHEET = "הסכמים מייצגים"


def sheet(wb, title, header, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(header)
    for j, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def put(ws, row, values, fill=None):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = WRAP
        c.border = BORDER
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)


def hyperlink(cell, target, text=None):
    cell.value = text if text is not None else target
    cell.hyperlink = target
    cell.font = Font(name=FONT, size=10, color="0B5CAD", underline="single")


def finish(ws, ncols, nrows):
    if nrows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(ncols), nrows + 1)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--research", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    r = normalise(json.load(io.open(a.research, encoding="utf-8")))
    blocking, flags, notes = check(r, stage="brief")
    if blocking:
        print("קובץ ה-Excel לא נבנה:")
        for b in blocking:
            print("  חוסם: " + b)
        sys.exit(1)

    wb = Workbook()
    wb.remove(wb.active)
    law_rows = {}

    if scope_has(r, "research"):
        header = ["מס\"ד", "דרג", "סוג המקור", "שם המקור", "גוף מפרסם", "מעמד", "תאריך פרסום מקורי",
                  "תאריך פרסום או תיקון אחרון", "סעיפים רלוונטיים", "תמצית הסעיפים", "המשמעות להסכם",
                  "אופן ההתייחסות בהסכם", "סוג השפעה", "סעיף חובה", "קישור", "סטטוס אימות"]
        ws = sheet(wb, LAW_SHEET, header, [6, 6, 16, 32, 20, 9, 12, 14, 16, 36, 36, 40, 12, 9, 34, 14])
        for i, src in enumerate(r.get("law") or []):
            n = i + 2
            rank = src.get("rank")
            status = src.get("status") or RANK_STATUS.get(rank, "")
            law_rows[src.get("row", i + 1)] = n
            put(ws, n, [src.get("row", i + 1), rank, src.get("source_kind") or NORM_RANKS.get(rank, ""),
                        src.get("instrument", ""), src.get("publisher", ""), status,
                        display_date(src.get("published_original")), display_date(src.get("published_last")),
                        src.get("sections", ""), src.get("sections_summary", ""), src.get("meaning", ""),
                        src.get("contract_treatment", ""), src.get("effect", ""),
                        "כן" if src.get("mandatory") else "לא", src.get("url", ""),
                        "אומת" if src.get("verified") else "טעון אימות"],
                STATUS_FILL.get(status))
            hyperlink(ws.cell(row=n, column=4), src["url"], src.get("instrument", ""))
            hyperlink(ws.cell(row=n, column=15), src["url"])
        finish(ws, len(header), len(r.get("law") or []))

        header = ["מזהה", "שם הסעיף", "הדרישה", "שורת מקור", "סעיף בחיקוק", "הצד החייב",
                  "נוסח מוצע בתמצית", "סטטוס", "אישור המשתמש"]
        ws = sheet(wb, MAND_SHEET, header, [8, 24, 40, 10, 14, 14, 44, 10, 16])
        for i, m in enumerate(r.get("mandatory_clauses") or []):
            n = i + 2
            put(ws, n, [m.get("clause_id", ""), m.get("title", ""), m.get("requirement", ""),
                        m.get("source_row", ""), m.get("source_section", ""), m.get("bound_party", ""),
                        m.get("proposed_text", ""), m.get("status", "מוצע"), ""])
            target = law_rows.get(m.get("source_row"))
            if target:
                hyperlink(ws.cell(row=n, column=4), "#'%s'!D%d" % (LAW_SHEET, target),
                          "שורה %s" % m.get("source_row"))
        finish(ws, len(header), len(r.get("mandatory_clauses") or []))

    if scope_has(r, "agreements"):
        header = ["מס'", "כותרת", "מפרסם", "סוג האתר", "דין חל", "כיוון", "נימוק לכיוון", "נושאים",
                  "מה ההסכם תורם", "סעיפי חובה חסרים", "תאריך פרסום", "שימוש", "קישור"]
        ws = sheet(wb, AGR_SHEET, header, [6, 30, 20, 14, 12, 14, 36, 30, 34, 16, 12, 22, 34])
        for i, ag in enumerate(r.get("agreements") or []):
            n = i + 2
            put(ws, n, [ag.get("no", i + 1), ag.get("title", ""), ag.get("publisher", ""),
                        ag.get("source_type", ""), ag.get("governing_law", ""), ag.get("orientation", ""),
                        ag.get("orientation_reason", ""), ", ".join(ag.get("topics") or []),
                        ag.get("contributes", ""), ", ".join(ag.get("missing_mandatory") or []),
                        display_date(ag.get("published")), ", ".join(flags["agreements"][i]),
                        ag.get("url", "")])
            hyperlink(ws.cell(row=n, column=2), ag["url"], ag.get("title", ""))
            hyperlink(ws.cell(row=n, column=13), ag["url"])
        finish(ws, len(header), len(r.get("agreements") or []))

    wb.save(a.out)
    print("קובץ ה-Excel: %s" % a.out)
    print("גיליונות: %s" % " · ".join(wb.sheetnames))
    for note in notes:
        print("  " + note)


if __name__ == "__main__":
    main()

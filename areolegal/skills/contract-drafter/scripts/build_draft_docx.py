#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""בונה את מסמך הטיוטה ב-Word ישירות מ-draft.json.

**זה המסלול היחיד להפקת טיוטה ב-Word.** אין לכתוב את המסמך ב-python-docx בכל
הרצה, ואין "לשפר" את העיצוב ביד. בנייה ידנית היא מה שגרם לטיוטה לצאת בברירות
המחדל של Word, בלי היררכיה ובלי כיווניות, והיא גם מה שגרם לשלב ההפקה להימשך
עשרות דקות: המודל נאלץ לפלוט סדרתית גם את נוסח ההסכם וגם את קוד הבנייה שלו.
כאן המודל פולט **נתונים בלבד**, והסקריפט עושה את המספור, ההיררכיה, הכיווניות
וסימון החוסרים.

הסקריפט אוכף כיווניות בעצמו. **אין להעביר את התוצר שלו ב-rtl-docx-enforcer:**
מעבר נוסף מייצר שכתוב מלא של המסמך בלי להוסיף דבר.

    python3 build_draft_docx.py --draft Areo-נתונים/draft-<slug>.json \\
        --profile Areo-נתונים/פרופיל-לקוח.json --out "<שם המסמך>.docx"

מבנה draft.json, במלואו. שדה חסר מדולג בשקט; אין שדה חובה מלבד clauses:
    meta{doc_type,title,language,classification,transaction_id,template{},
         parties[{defined_as,name,reg,address,role}],place_date,generated_at}
    recitals[str]  definitions[{term,text}]
    clauses[{number,heading,text,children[...]}]
    annexes[{label,title,text|clauses[]}]
    signatures{note,blocks[{name,by,title}]}
"""
import argparse
import datetime
import io
import json
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------- פלטה: מסמך משפטי, לא דוח. מאופקת בכוונה ----------
INK = "16191F"
INK_SOFT = "5B6472"
LINE = "C9CFD6"
NAVY = "12314F"
GAP_BG = "FFF2A8"      # רקע לציון חוסר, כדי שלא ייחתם מסמך עם [להשלים]
FONT = "Arial"

RTL = True             # נקבע מ-meta.language בזמן הבנייה

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "numPr", "pBdr",
             "shd", "tabs", "bidi", "spacing", "ind", "contextualSpacing", "jc",
             "outlineLvl", "rPr"]

GAP_RX = re.compile(r"(\[[^\[\]]{0,80}\])")


# ---------- כלי XML ----------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def _local(tag):
    return tag.split("}")[-1]


# אותו כלל סדר חל על w:tblPr ועל w:sectPr, ושם הוא לא נאכף עד כה: bidiVisual
# שצורף בסוף tblPr נחת אחרי tblBorders ו-tblLook, ו-bidi בסוף sectPr נחת אחרי
# docGrid. Word פותח בכל זאת, אבל הקובץ אינו תקף מול הסכמה.
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
               "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
               "tblDescription", "tblPrChange"]
SECTPR_ORDER = ["footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
                "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
                "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
                "printerSettings", "sectPrChange"]


def put_ordered(parent, el, order):
    """מכניס אלמנט למקומו לפי הסכמה, ומחליף אלמנט קיים מאותו סוג."""
    name = _local(el.tag)
    old = parent.find(qn("w:" + name))
    if old is not None:
        parent.remove(old)
    idx = order.index(name) if name in order else 99
    for child in parent:
        c = _local(child.tag)
        if (order.index(c) if c in order else 99) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def ppr_put(p, el):
    """סדר הילדים ב-w:pPr קבוע בסכמת OOXML. אלמנט שמצורף בסוף מתעלמים ממנו בשקט."""
    ppr = p._p.get_or_add_pPr()
    name = _local(el.tag)
    old = ppr.find(qn("w:" + name))
    if old is not None:
        ppr.remove(old)
    idx = PPR_ORDER.index(name) if name in PPR_ORDER else 99
    for child in ppr:
        c = _local(child.tag)
        if (PPR_ORDER.index(c) if c in PPR_ORDER else 99) > idx:
            child.addprevious(el)
            return el
    ppr.append(el)
    return el


def dir_p(p):
    if RTL:
        ppr_put(p, _el("w:bidi"))


def run(p, text, size=11, bold=False, color=INK, italic=False, fill=None, underline=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.underline = underline
    r.font.color.rgb = RGBColor.from_string(color)
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    for k in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + k), FONT)
    if fill:
        rpr.append(_el("w:shd", val="clear", color="auto", fill=fill))
    if RTL:
        rpr.append(_el("w:rtl"))
    return r


GAPS = []


LINK = "0B5CAD"


def link(p, url, text=None, size=10, bold=False):
    """היפר-קישור לחיץ: w:hyperlink עם r:id שנרשם כקשר חיצוני של המסמך.
    כתובת שנכתבת כטקסט רגיל אינה לחיצה ב-Word, ולכן אינה נחשבת קישור."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    url = (url or "").strip()
    if not url:
        return None
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = _el("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = _el("w:r")
    rpr = _el("w:rPr")
    rf = _el("w:rFonts")
    for k in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + k), FONT)
    rpr.append(rf)
    if bold:
        rpr.append(_el("w:b"))
    rpr.append(_el("w:color", val=LINK))
    rpr.append(_el("w:sz", val=str(int(size * 2))))
    rpr.append(_el("w:u", val="single"))
    if RTL:
        rpr.append(_el("w:rtl"))
    r.append(rpr)
    t = _el("w:t")
    t.text = text or url
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)
    return h


def cell_link(cell, url, text=None, size=9.5):
    p = cell.paragraphs[0]
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
    dir_p(p)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if url:
        link(p, url, text or url, size)
    else:
        run(p, text or NOT_FOUND_TEXT, size, False, INK_SOFT)


NOT_FOUND_TEXT = "[לא אותר]"


def rich(p, text, size=11, bold=False, color=INK):
    """כותב טקסט ומסמן כל ציין חוסר ברקע, כדי שלא ייעלם בקריאה."""
    for part in GAP_RX.split(s(text)):
        if not part:
            continue
        if GAP_RX.fullmatch(part):
            GAPS.append(part)
            run(p, part, size, True, INK, fill=GAP_BG)
        else:
            run(p, part, size, bold, color)
    return p


def para(doc, text="", size=11, bold=False, color=INK, before=0, after=5, line=1.2,
         container=None, align=None, indent=0.0, italic=False):
    p = (container or doc).add_paragraph()
    dir_p(p)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        pf.left_indent = Cm(indent) if not RTL else None
        pf.right_indent = Cm(indent) if RTL else None
    if text:
        run(p, text, size, bold, color, italic)
    return p


def table_borders(tbl, color=LINE, visible=True):
    tpr = tbl._tbl.tblPr
    b = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el("w:" + edge, val="single" if visible else "none",
                     sz="4", space="0", color=color))
    put_ordered(tpr, b, TBLPR_ORDER)
    if RTL:
        put_ordered(tpr, _el("w:bidiVisual"), TBLPR_ORDER)


def cell_text(cell, text, bold=False, size=10.5, color=INK, align=None):
    p = cell.paragraphs[0]
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)           # cell.text = "" leaves a run with no rPr, hence no w:rtl
    dir_p(p)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if align is not None:
        p.alignment = align
    # split BEFORE s(): it normalises whitespace and would swallow the newlines
    lines = str(text).split("\n") if text else [""]
    for i, line in enumerate(lines):
        if i:
            run(p, "", size, bold, color)._r.append(_el("w:br"))
        rich(p, line, size, bold, color)


# ---------- ניקוי ----------
def s(v, limit=None):
    """מקף ארוך אסור בכל תוצר, וגם רווחים כפולים שנוצרים בהרכבת נוסח."""
    t = " ".join(str(v).split()) if v not in (None, "") else ""
    t = t.replace("—", ", ").replace("–", "-").replace(" ,", ",")
    return t[:limit] if limit else t


def produced_on(meta):
    raw = s(meta.get("generated_at"))[:10]
    if len(raw) == 10 and raw[4] == "-":
        y, mo, d = raw.split("-")
        return "%s.%s.%s" % (d, mo, y)
    return datetime.date.today().strftime("%d.%m.%Y")


def txt(node, *keys):
    for k in keys:
        if node.get(k):
            return s(node[k])
    return ""


# ---------- בלוקי המסמך ----------
def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:cs"), FONT)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.4)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.0)
        if RTL:
            put_ordered(sec._sectPr, _el("w:bidi"), SECTPR_ORDER)


def title_block(doc, m):
    t = txt(m, "title", "doc_type") or "טיוטה"
    p = para(doc, before=0, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, t, 19, True, NAVY)
    sub = " · ".join([x for x in (txt(m, "doc_type") if txt(m, "title") else "",
                                  txt(m, "place_date")) if x])
    if sub:
        p = para(doc, before=0, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        run(p, sub, 10.5, False, INK_SOFT)
    stamp = "טיוטה · הופק בתאריך %s" % produced_on(m)
    cls = txt(m, "classification")
    if cls:
        stamp += " · " + cls
    p = para(doc, before=0, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, stamp, 9.5, False, INK_SOFT)
    notice = txt(m, "notice")
    if notice:
        box = para(doc, before=0, after=10)
        ppr_put(box, _el("w:shd", val="clear", color="auto", fill="F5F7F9"))
        run(box, notice, 10, True, NAVY)
    rule = para(doc, before=0, after=14)
    ppr = rule._p.get_or_add_pPr()
    b = ppr_put(rule, _el("w:pBdr"))
    b.append(_el("w:top", val="single", sz="8", space="6", color=LINE))


def parties(doc, m):
    ps = m.get("parties") or []
    if not ps:
        return
    for i, party in enumerate(ps):
        p = para(doc, before=2, after=2)
        name = txt(party, "name") or "[להשלים]"
        head = name
        for extra in (txt(party, "reg"), txt(party, "address")):
            if extra:
                head += ", " + extra
        rich(p, head, 11, False)
        defined = txt(party, "defined_as") or txt(party, "role")
        if defined:
            run(p, ' (להלן: "%s")' % defined, 11, True)
        if i < len(ps) - 1:
            para(doc, "מצד אחד;" if i == 0 and len(ps) == 2 else "מצד אחד;",
                 10, False, INK_SOFT, after=6, align=WD_ALIGN_PARAGRAPH.LEFT
                 if RTL else WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, "מצד שני;", 10, False, INK_SOFT, after=12,
         align=WD_ALIGN_PARAGRAPH.LEFT if RTL else WD_ALIGN_PARAGRAPH.RIGHT)


def recitals(doc, rs):
    if not rs:
        return
    for r in rs:
        body = s(r)
        p = para(doc, before=0, after=5, indent=0.6)
        if not body.startswith("הואיל") and RTL:
            run(p, "הואיל ו", 11, True)
        rich(p, body, 11)
    para(doc, "לפיכך הוסכם, הותנה והוצהר בין הצדדים כדלקמן:", 11, True, INK,
         before=8, after=12)


def definitions(doc, defs):
    if not defs:
        return
    h1(doc, "הגדרות")
    for d in defs:
        p = para(doc, before=2, after=4, indent=0.6)
        run(p, '"%s" ' % txt(d, "term"), 11, True)
        run(p, "- ", 11, False, INK_SOFT)
        rich(p, txt(d, "text", "definition"), 11)


def h1(doc, text, num=""):
    p = para(doc, before=16, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if num:
        run(p, "%s. " % num, 13, True, NAVY)
    run(p, s(text), 13, True, NAVY)
    return p


def clause(doc, c, depth=0):
    """סעיף ותתי-סעיפיו. המספר נכתב לתוך הטקסט ולא כמספור אוטומטי של Word,
    כדי שבקרת השלמות תמצא אותו ושהכיווניות לא תהפוך את הספרה והנקודה."""
    num = s(c.get("number"))
    head = txt(c, "heading", "title")
    body = txt(c, "text", "body")
    if c.get("needs_case_law") and body:
        # the law sources in state 4 exclude case law, so a clause whose validity rests
        # on it (liability caps, non-compete, liquidated damages) is marked for the lawyer
        body += " [טעון בדיקת פסיקה]"
    if depth == 0:
        if head:
            h1(doc, head, num)
            if body:
                rich(para(doc, before=0, after=5, indent=0.6), body, 11)
        else:
            p = para(doc, before=8, after=5)
            if num:
                run(p, "%s. " % num, 11, True)
            rich(p, body, 11)
    else:
        p = para(doc, before=0, after=4, indent=0.6 + 0.6 * depth)
        if num:
            run(p, "%s " % num, 11, True)
        if head:
            run(p, s(head) + ". ", 11, True)
        rich(p, body, 11)
    for ch in (c.get("children") or c.get("subclauses") or []):
        clause(doc, ch, depth + 1)


def page_break(doc):
    p = para(doc, before=0, after=0)
    run(p, "")._r.append(_el("w:br", type="page"))
    return p


def annexes(doc, ans):
    if not ans:
        return
    for a in ans:
        page_break(doc)
        label = txt(a, "label") or "נספח"
        title = txt(a, "title")
        p = para(doc, before=0, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        run(p, label + (" - " + title if title else ""), 15, True, NAVY)
        if a.get("text"):
            rich(para(doc, before=0, after=5), txt(a, "text"), 11)
        for c in (a.get("clauses") or []):
            clause(doc, c)


def signatures(doc, sig, m):
    blocks = (sig or {}).get("blocks") or []
    if not blocks:
        blocks = [{"name": txt(p, "defined_as", "name")} for p in (m.get("parties") or [])]
    if not blocks:
        return
    para(doc, s((sig or {}).get("note")) or "ולראיה באו הצדדים על החתום:",
         11, True, INK, before=22, after=14)
    tbl = doc.add_table(rows=2, cols=len(blocks))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(tbl, LINE, visible=False)
    for i, b in enumerate(blocks):
        cell_text(tbl.rows[0].cells[i], "_______________________", False, 11,
                  INK_SOFT, WD_ALIGN_PARAGRAPH.CENTER)
        lines = [txt(b, "name")]
        if txt(b, "by"):
            lines.append("על ידי: " + txt(b, "by"))
        if txt(b, "title"):
            lines.append(txt(b, "title"))
        cell_text(tbl.rows[1].cells[i], "\n".join([x for x in lines if x]),
                  False, 10, INK, WD_ALIGN_PARAGRAPH.CENTER)


def clause_sources(c):
    """"MAND-M1, REG-2" או רשימה: מזהי המקור של סעיף."""
    v = c.get("source")
    if isinstance(v, (list, tuple)):
        items = v
    else:
        items = [x for x in re.split(r"[,;]", str(v or ""))]
    return [x.strip() for x in items if str(x).strip()]


def walk_clauses(cs, out=None):
    out = [] if out is None else out
    for c in cs or []:
        out.append(c)
        walk_clauses(c.get("children") or c.get("subclauses"), out)
    return out


def all_clauses(d):
    out = walk_clauses(d.get("clauses"))
    for a in d.get("annexes") or []:
        walk_clauses(a.get("clauses"), out)
    return out


def source_link(sid, research):
    """טקסט וקישור למזהה מקור: REG-<שורה>, MAND-<מזהה>, AGR-<מספר>."""
    from research_sources import approved_agreements, normalise
    r = normalise(research or {})
    kind, _, key = sid.partition("-")
    if kind == "REG":
        for i, x in enumerate(r.get("law") or []):
            if str(x.get("row", i + 1)) == key:
                return "%s (%s)" % (s(x.get("instrument")), s(x.get("sections"))), x.get("url")
    elif kind == "AGR":
        for x in approved_agreements(r):
            if str(x.get("no")) == key:
                return s(x.get("title")), x.get("url")
    elif kind == "MAND":
        rows = {str(x.get("row", i + 1)): x for i, x in enumerate(r.get("law") or [])}
        for m in r.get("mandatory_clauses") or []:
            if m.get("clause_id") == key:
                src = rows.get(str(m.get("source_row")), {})
                return "סעיף חובה %s: %s" % (key, s(m.get("title"))), src.get("url")
    return sid, None


def sources_annex(doc, d, research):
    """נספח המקורות של מצב 4. פנימי: מוסר לפני שליחה החוצה. כל מקור לחיץ."""
    if not research:
        return
    from research_sources import NORM_RANKS, approved_agreements, normalise, scope_has
    r = normalise(research)
    law = (r.get("law") or []) if scope_has(r, "research") else []
    ags = approved_agreements(r) if scope_has(r, "agreements") else []
    page_break(doc)
    p = para(doc, before=0, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "נספח מקורות", 15, True, NAVY)
    para(doc, "פנימי. יש להסיר לפני שליחה לצד שכנגד.", 9.5, True, INK_SOFT, after=10,
         align=WD_ALIGN_PARAGRAPH.CENTER)

    def grid(header, rows, link_col):
        tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(tbl)
        for jj, h in enumerate(header):
            cell_text(tbl.rows[0].cells[jj], h, True, 9.5, NAVY)
        for ii, row in enumerate(rows):
            for jj, v in enumerate(row):
                if jj == link_col:
                    cell_link(tbl.rows[ii + 1].cells[jj], v[1], v[0])
                else:
                    cell_text(tbl.rows[ii + 1].cells[jj], v, False, 9.5)
        para(doc, after=6)

    if law:
        h1(doc, "הדין החל")
        grid(["מזהה", "דרג", "מקור", "סעיפים", "סטטוס"],
             [["REG-%s" % x.get("row", k + 1), NORM_RANKS.get(x.get("rank"), ""),
               (s(x.get("instrument")), x.get("url")), s(x.get("sections")),
               "אומת" if x.get("verified") else "טעון אימות"] for k, x in enumerate(law)], 2)
    if ags:
        h1(doc, "הסכמים מייצגים שאושרו")
        grid(["מזהה", "הסכם", "מפרסם", "כיוון"],
             [["AGR-%s" % x.get("no"), (s(x.get("title")), x.get("url")), s(x.get("publisher")),
               s(x.get("orientation"))] for x in ags], 1)


def footer(doc, m):
    bits = [txt(m, "doc_type") or "טיוטה"]
    if txt(m, "classification"):
        bits.append(txt(m, "classification"))
    bits.append(produced_on(m))
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        dir_p(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, " · ".join(bits) + " · עמוד ", 8.5, False, INK_SOFT)
        p._p.append(_el("w:fldSimple", instr="PAGE"))


# ---------- הרכבה ----------
class State4Refused(Exception):
    pass


def state4_gate(d, research):
    """מצב 4 מזוהה לפי notice, sources או קובץ מחקר. כך אי אפשר לעקוף את הבדיקה
    פשוט בהשמטת דגל. הסקריפט בודק שהאישורים **נרשמו**; שהמשתמש אכן לחץ נאכף בסקיל."""
    from research_sources import approved_mandatory, check, normalise, scope_has
    m = d.get("meta") or {}
    if not (m.get("notice") or d.get("sources") or research):
        return None
    if not research:
        raise State4Refused(["טיוטת מצב 4 (notice או sources) נבנית רק עם --research."])
    blocking, _, _ = check(research, stage="draft")
    r = normalise(research)
    if scope_has(r, "research"):
        placed = set()
        for c in all_clauses(d):
            placed.update(clause_sources(c))
        for mc in approved_mandatory(r):
            if "MAND-%s" % mc.get("clause_id") not in placed:
                blocking.append("סעיף חובה מאושר %s (%s) לא שובץ באף סעיף בטיוטה."
                                % (mc.get("clause_id"), mc.get("title") or ""))
    if blocking:
        raise State4Refused(blocking)
    return r


def build(d, prof, out, research=None, cover=None):
    global RTL
    m = d.get("meta") or {}
    r4 = state4_gate(d, research)
    lang = s(m.get("language")) or s((prof or {}).get("deliverable_language")) or "he"
    RTL = lang.lower() not in ("en", "english", "אנגלית")

    del GAPS[:]
    doc = Document()
    base(doc)
    title_block(doc, m)
    parties(doc, m)
    recitals(doc, d.get("recitals"))
    definitions(doc, d.get("definitions"))
    for c in (d.get("clauses") or []):
        clause(doc, c)
    annexes(doc, d.get("annexes"))
    signatures(doc, d.get("signatures"), m)
    sources_annex(doc, d, r4)
    footer(doc, m)
    doc.save(out)
    if r4 is not None and cover:
        write_cover4(cover, d, r4)

    # counted from what was actually written, not from the input: a placeholder
    # that never reached the page must not be reported as marked.
    gaps = len(GAPS)
    return {"out": out, "rtl": RTL, "clauses": len(d.get("clauses") or []),
            "definitions": len(d.get("definitions") or []),
            "annexes": len(d.get("annexes") or []), "placeholders": gaps}


def write_cover4(path, d, r):
    """דף הליווי של מצב 4: מטריצת עקיבות, סעיפי חובה שהוסרו, פערים, והצעה להערכת סיכונים."""
    from research_sources import SCOPES, normalise, scope_has
    r = normalise(r)
    doc = Document()
    base(doc)
    m = d.get("meta") or {}
    run(para(doc, after=2), "דף ליווי: " + (txt(m, "title", "doc_type") or "טיוטה"), 16, True, NAVY)
    run(para(doc, after=10), "טיוטה · הופק בתאריך %s · מקור: מצב 4, %s" % (produced_on(m), SCOPES[r["scope"]]),
        9.5, False, INK_SOFT)

    h1(doc, "מטריצת עקיבות")
    rows = [c for c in all_clauses(d) if c.get("number")]
    tbl = doc.add_table(rows=1 + len(rows), cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(tbl)
    for jj, h in enumerate(["סעיף", "נושא", "מקור", "סעיף חובה", "עמדת פתיחה"]):
        cell_text(tbl.rows[0].cells[jj], h, True, 9.5, NAVY)
    for ii, c in enumerate(rows):
        sids = clause_sources(c)
        cells = tbl.rows[ii + 1].cells
        cell_text(cells[0], s(c.get("number")), False, 9.5)
        cell_text(cells[1], txt(c, "heading") or txt(c, "text")[:40], False, 9.5)
        if sids:
            label, url = source_link(sids[0], r)
            cell_link(cells[2], url, ", ".join(sids)) if url else cell_text(cells[2], ", ".join(sids), False, 9.5)
        else:
            cell_text(cells[2], "ניסוח עצמאי", False, 9.5, INK_SOFT)
        cell_text(cells[3], "כן" if any(x.startswith("MAND-") for x in sids) else "לא", False, 9.5)
        cell_text(cells[4], "כן" if c.get("opening_position") else "", False, 9.5)
    para(doc, after=6)

    ap_ = r.get("approvals") or {}
    removed = ap_.get("removed_clauses") or []
    if removed:
        h1(doc, "סעיפי חובה שהוסרו באישור המשתמש")
        for x in removed:
            if isinstance(x, dict):
                line = "%s: %s. החשיפה: %s" % (x.get("clause_id", ""), x.get("reason", ""), x.get("exposure", "[להשלים]"))
            else:
                line = str(x)
            rich(para(doc, after=3, indent=0.5), line, 10.5)

    h1(doc, "פערים ומגבלות")
    if not scope_has(r, "research"):
        para(doc, "לא בוצע מחקר רגולטורי: נבחר איסוף הסכמים בלבד. הטיוטה אינה נשענת על מיפוי של הדין החל.",
             10.5, True, INK, after=3)
    else:
        para(doc, "המחקר מוגבל למקורות רשמיים (gov.il, אתר הכנסת, נבו בנתיבי חקיקה ואתרים ממשלתיים זרים "
                  "מרשימה סגורה), ואינו מכסה פסיקה. הנחיה שלא פורסמה במקורות אלה, או שאינה נגישה לשליפה, "
                  "אינה בו.", 10.5, False, INK, after=3)
        for x in [x for x in (r.get("law") or []) if not x.get("verified")]:
            p = para(doc, after=3, indent=0.5)
            run(p, "טעון אימות: ", 10.5, True)
            link(p, x.get("url"), s(x.get("instrument")), 10.5)
    marked = [c for c in all_clauses(d) if c.get("needs_case_law")]
    if marked:
        para(doc, "טעון בדיקת פסיקה: סעיפים " + ", ".join(s(c.get("number")) for c in marked), 10.5, False, INK, after=3)
    para(doc, "להריץ על הטיוטה הערכת סיכונים לפני שהיא יוצאת?", 10.5, True, INK, before=14)
    doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--draft", required=True)
    ap.add_argument("--profile")
    ap.add_argument("--out", required=True)
    ap.add_argument("--research", help="מצב 4: קובץ המחקר שאושר. בלעדיו טיוטת מצב 4 אינה נבנית")
    ap.add_argument("--cover", help="מצב 4: דף ליווי עם מטריצת עקיבות")
    a = ap.parse_args()
    d = json.load(io.open(a.draft, encoding="utf-8"))
    prof = {}
    if a.profile:
        try:
            prof = json.load(io.open(a.profile, encoding="utf-8"))
        except (OSError, ValueError):
            prof = {}
    research = json.load(io.open(a.research, encoding="utf-8")) if a.research else None
    try:
        r = build(d, prof, a.out, research, a.cover)
    except State4Refused as e:
        print("הטיוטה לא נבנתה:")
        for b in e.args[0]:
            print("  חוסם: " + b)
        sys.exit(1)
    print("הטיוטה נבנתה: %s" % r["out"])
    print("סעיפים ראשיים: %d · הגדרות: %d · נספחים: %d · כיווניות: %s"
          % (r["clauses"], r["definitions"], r["annexes"], "RTL" if r["rtl"] else "LTR"))
    if r["placeholders"]:
        print("מצייני חוסר שסומנו ברקע: %d. אין לסווג ReadyForExternal כל עוד הם פתוחים."
              % r["placeholders"])


if __name__ == "__main__":
    main()

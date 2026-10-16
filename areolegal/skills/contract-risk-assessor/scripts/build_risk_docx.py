#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""גרסת ה-Word של דוח הסיכונים, רק לבקשת המשתמש, באותו מבנה ובאותו סדר בדיוק.

`methodology.md`, "תוצר": ה-Word נבנה מאותו מרשם, באותו מבנה ובאותו סדר כמו
הדוח, בלי פרקים נוספים, ולעולם אינו ארוך מהדוח (הערות מעריך הסיכונים 3.6,
שינויים 1.15, 11.13, 16.5, 18.10 ו-21.3):

1. **עמוד 1, התמונה במבט אחד:** שורת הכותרת, השורה התחתונה, האם אפשר לחתום כמו
   שהוא, "לתיקון לפני שליחת התגובה" כשיש, טבלת העדשות, שלושת המונים, המועדים
   הקריטיים ושורת חקר הישות.
2. **פרק 1, הסיכונים:** טבלה אחת לפי חומרה. בהסכם הדדי, הצד בשם הסיכון.
3. **פרק 2, סתירות בין ההסכם לנספחים**, ומה גובר.
4. **פרק 3, צירופים ושרשראות.**
5. **פרק 4, הסיכונים הקריטיים בפירוט.**
6. **פרק 5, מה לא נבדק ומה חסר**, וטבלת המקורות "מה נבדק ברשת ומתי", עם קישור
   אמיתי לכל מקור והכתובת גם כטקסט, להדפסה.
7. **נספח: ציטוטים מילוליים.**

**סימון החיסיון** ("טיוטה, טעונה אישור" והסימון של עורך הדין) בכותרת העליונה
ובכותרת התחתונה של כל עמוד, באותו נוסח כמו בדוח. הקובץ נשמר בתיקיית העסקה,
לצד המרשם, ושמו כולל את התאריך; קובץ קיים אינו נדרס.

    python3 scripts/build_risk_docx.py --data "<תיקיית העסקה>/risk-register.json"

קוד יציאה: **0** נבנה, **2** קלט חסר, **3** בדיקת הכיווניות נכשלה.
"""
import argparse
import datetime as _dt
import io
import json
import os
import re
import sys

import areo_lib  # noqa: F401  מוסיף את areolegal/lib לנתיב הייבוא

from docx import Document  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

from common import paths, rtl  # noqa: E402
from risk.render_check import band_text_from_register  # noqa: E402

LENS_HE = {"legal": "משפטי-חוזי", "financial": "כספי", "operational": "תפעולי", "technology": "טכנולוגי וסייבר",
           "regulatory": "רגולציה ופרטיות", "reputational": "מוניטין", "insurance": "ביטוח",
           "force_majeure": "כוח עליון וגיאופוליטי"}
LEVEL_HE = {"critical": "קריטי", "high": "גבוה", "medium": "בינוני", "low": "נמוך"}
LEVEL_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3}
SIGN_HE = {"yes": "אפשר לחתום כמו שהוא", "with_fixes": "אפשר לחתום אחרי התיקונים", "no": "לא מומלץ לחתום כמו שהוא"}
SIDE_HE = {"as_obligor": "כמי שמתחייב", "as_beneficiary": "כמי שהסעיף מגן עליו", "both": "בשני התפקידים",
           "as_discloser": "כמי שהסעיף מגן עליו", "as_recipient": "כמי שמתחייב"}
STATUS_HE = {"findings": "נמצאו ממצאים", "none_found": "נבדק ולא נמצא", "not_checked": "לא נבדק"}


def load(path):
    if not os.path.exists(path):
        print("המרשם אינו קיים: %s" % path)
        raise SystemExit(2)
    with io.open(path, encoding="utf-8") as fh:
        return json.load(fh)


def h(doc, text, size=14):
    p = rtl.add_paragraph(doc, text, bold=True, size=size)
    return p


def para(doc, text, size=10.5, bold=False):
    return rtl.add_paragraph(doc, str(text or ""), bold=bold, size=size)


def table(doc, head, rows):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, x in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = x
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v if v is not None else "")
    rtl.set_table_rtl(t)
    return t


def hyperlink(paragraph, url, text):
    """קישור אמיתי (w:hyperlink עם יחס חיצוני), והכתובת עצמה גם כטקסט, כדי שתודפס."""
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    tail = paragraph.add_run(" (%s)" % url)
    tail.font.size = Pt(8)
    rtl.set_run_ltr(tail)
    return paragraph


def band(doc, text):
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            p = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
            r = p.add_run(text)
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)


def quotes_index(reg):
    """כל ציטוט מקבל מספר אחד בנספח, וכל הפניה בדוח היא למספר הזה."""
    seen, out = {}, []
    def add(a):
        if not a or not a.get("quote"):
            return None
        key = (a.get("document"), a.get("clause"), a.get("quote"))
        if key not in seen:
            seen[key] = len(out) + 1
            out.append(a)
        return seen[key]
    for r in reg.get("risks") or []:
        r["_q"] = [add(a) for a in r.get("anchors") or []]
    for d in reg.get("document_defects") or []:
        d["_q"] = [add(a) for a in d.get("locations") or []]
    for c in reg.get("dangerous_combinations") or []:
        c["_q"] = [add(a) for a in c.get("anchors") or []]
    return out


def refs(nums):
    return " ".join("[%d]" % n for n in nums or [] if n)


def build(reg, out):
    m, s = reg.get("meta") or {}, reg.get("summary") or {}
    quotes = quotes_index(reg)
    risks = sorted(reg.get("risks") or [], key=lambda r: (LEVEL_ORDER.get(r.get("severity_level"), 9),
                                                          r.get("risk_no") or 0))
    mutual = "הדדי" in str(m.get("position") or "") or m.get("mutual_primary_role")
    doc = Document()
    band(doc, band_text_from_register(reg))

    h(doc, "Areo Workspace | סיכונים", 16)
    h(doc, m.get("scope_label") or "", 13)
    h(doc, "התמונה במבט אחד")
    cp = (m.get("counterparty") or {}).get("name") if isinstance(m.get("counterparty"), dict) else m.get("counterparty")
    for label, val in (("הצד שכנגד", cp), ("העסקה", s.get("deal_one_line")), ("הפוזיציה שלנו", m.get("position")),
                       ("מה נסרק", ", ".join(m.get("annexes_reviewed") or []) if isinstance(m.get("annexes_reviewed"),
                                                                                         list) else m.get("annexes_reviewed")),
                       ("מה לא נמסר", ", ".join(m.get("annexes_missing") or []) if isinstance(m.get("annexes_missing"),
                                                                                           list) else m.get("annexes_missing")),
                       ("הנושא שביקשתם", m.get("user_focus"))):
        if val:
            para(doc, "%s: %s" % (label, val))
    h(doc, "השורה התחתונה", 12)
    for line in s.get("bottom_line") or []:
        para(doc, line)
    if s.get("sign_as_is"):
        para(doc, SIGN_HE.get(s["sign_as_is"], s["sign_as_is"]), bold=True)
    fixes = [f for f in reg.get("pre_send_fixes") or [] if f.get("status", "open") == "open"]
    if fixes:
        h(doc, "לתיקון לפני שליחת התגובה", 12)
        for f in fixes:
            para(doc, "%s (%s). סיכונים: %s" % (f.get("gap"), f.get("where"),
                                                ", ".join(str(x).replace("RSK-0", "").replace("RSK-", "")
                                                          for x in f.get("risk_ids") or [])))
    h(doc, "העדשות", 12)
    table(doc, ("עדשה", "מצב", "סיכונים"), [(LENS_HE.get(c.get("lens"), c.get("lens")),
                                             STATUS_HE.get(c.get("status"), c.get("status")),
                                             len(c.get("risk_ids") or [])) for c in reg.get("coverage") or []])
    para(doc, "סיכונים: %d   סתירות: %d   צירופים: %d" % (len(risks), len(reg.get("document_defects") or []),
                                                          len(reg.get("dangerous_combinations") or [])), bold=True)
    if s.get("key_dates"):
        h(doc, "מועדים קריטיים", 12)
        table(doc, ("מתי", "מה קורה", "סעיף", "מה אם חולף"),
              [(k.get("date") or k.get("window"), k.get("what"), k.get("clause"), k.get("consequence"))
               for k in s["key_dates"][:5]])
    er = reg.get("entity_research") or {}
    para(doc, er.get("note") or ("בדיקת הצד שכנגד לא בוצעה." if er.get("status") != "done" else ""))

    h(doc, "1. הסיכונים")
    rows = []
    for r in risks:
        name = r.get("name") or ""
        if mutual and r.get("side"):
            name = "%s: %s" % (SIDE_HE.get(r["side"], r["side"]), name)
        a = (r.get("anchors") or [{}])[0]
        rows.append((r.get("risk_no"), LENS_HE.get(r.get("primary_lens"), r.get("primary_lens")), name,
                     "%s %s" % (a.get("clause") or "", refs(r.get("_q"))), LEVEL_HE.get(r.get("severity_level"), ""),
                     r.get("treatment_direction")))
    table(doc, ("מספר", "עדשה", "הסיכון בשורה אחת", "סעיף", "חומרה", "כיוון טיפול"), rows)

    h(doc, "2. סתירות בין ההסכם לנספחים")
    defects = reg.get("document_defects") or []
    if defects:
        table(doc, ("נושא", "למה מסוכן", "מה גובר", "התיקון"),
              [(d.get("topic") or d.get("title"), "%s %s" % (d.get("risk_created") or "", refs(d.get("_q"))),
                d.get("prevails") or d.get("what_prevails") or "", d.get("fix") or d.get("proposed_fix") or "")
               for d in defects])
    else:
        para(doc, "לא נמצאו סתירות.")

    h(doc, "3. צירופים ושרשראות")
    combos = reg.get("dangerous_combinations") or []
    for c in combos[:5]:
        para(doc, "%s %s" % (c.get("why_together") or "", refs(c.get("_q"))))
        if c.get("breaking_fix"):
            para(doc, "התיקון ששובר את הצירוף: %s" % c["breaking_fix"], size=10)
    if not combos:
        para(doc, "לא נמצאו צירופים.")

    h(doc, "4. הסיכונים הקריטיים בפירוט")
    crit = [r for r in risks if r.get("severity_level") == "critical"][:5]
    for r in crit:
        para(doc, "%s. %s" % (r.get("risk_no"), r.get("name")), bold=True)
        para(doc, r.get("critical_detail") or r.get("description"))
    if not crit:
        para(doc, "לא דורג סיכון קריטי.")

    h(doc, "5. מה לא נבדק ומה חסר")
    for x in s.get("not_checked") or []:
        para(doc, "%s: %s (%s)" % (x.get("item"), x.get("why"), x.get("who")) if isinstance(x, dict) else x)
    for x in s.get("assumptions") or []:
        para(doc, "הנחה לאישור: %s" % ((x.get("text") or x.get("assumption") or "") if isinstance(x, dict) else x))
    src = (reg.get("research") or {}).get("sources_checked") or []
    if src:
        h(doc, "מה נבדק ברשת ומתי", 12)
        for n, x in enumerate(src, 1):
            p = para(doc, "%d. %s: %s. " % (n, x.get("source"), x.get("taken") or x.get("subject") or ""), size=10)
            if x.get("url"):
                hyperlink(p, x["url"], "פתיחה")
            p.add_run(" %s" % (x.get("date") or ""))
        para(doc, "הבדיקה אינה תחליף לבדיקת נאותות פיננסית מלאה.", size=9)

    h(doc, "נספח: ציטוטים מילוליים")
    table(doc, ("מספר", "מסמך וסעיף", "הנוסח המילולי"),
          [(n, "%s, %s" % (a.get("document"), a.get("clause")), a.get("quote")) for n, a in enumerate(quotes, 1)])
    rtl.enforce_document(doc, author=m.get("prepared_for_counsel") or "")
    doc.save(out)


def main(argv=None):
    sys.dont_write_bytecode = True
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--data", required=True, help="risk-register.json, בתיקיית העסקה")
    ap.add_argument("--date", help="DD.MM.YYYY; בלעדיו: היום")
    a = ap.parse_args(argv)
    reg = load(a.data)
    date = a.date or _dt.date.today().strftime("%d.%m.%Y")
    label = re.sub(r'[\\/:*?"<>|]', "", re.sub(r",\s*\d{2}\.\d{2}\.\d{4}$", "",
                                                (reg.get("meta") or {}).get("scope_label") or "הסכם"))
    out = paths.unique_path(os.path.join(os.path.dirname(os.path.abspath(a.data)),
                                         "%s - דוח סיכונים - %s.docx" % (label, date)))
    paths.check_write_path(out)
    build(reg, out)
    problems = rtl.audit_docx(out, run_date=_dt.date.today())
    print("נבנה: %s" % out)
    for p in problems:
        print("  בדיקת כיווניות: %s" % p)
    return 3 if problems else 0


if __name__ == "__main__":
    sys.exit(main())

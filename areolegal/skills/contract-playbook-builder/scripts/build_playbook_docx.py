#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""בונה את מסמך הפלייבוק ב-Word, בגרסה המקוצרת או המלאה, ישירות מ-playbook-model.json.

**זה המסלול היחיד להפקת הפלייבוק ב-Word.** אין לכתוב את המסמך מחדש בכל הרצה:
כך נוצרו מסמכים בלי שום עיצוב, טקסט שחור על לבן בברירות המחדל של Word, בלי צבעי
הרמזור שהם שפת המוצר.

העיצוב: Arial (מותקן בכל מחשב ומכסה עברית), כותרות בכחול כהה, ירוק Areo להדגשה,
ושלושת המסלולים בירוק, בענבר ובאדום עם קו צבע ורקע עדין. כריכה, סימון חיסיון,
כותרת תחתונה עם מספר עמוד, וכיווניות מימין לשמאל בכל פסקה, ריצה, טבלה וסעיף.

    python3 build_playbook_docx.py --model Areo-נתונים/playbook-model.json \\
        --profile Areo-נתונים/פרופיל-לקוח.json --out "<שם>.docx"

המסמך נבנה כך שיעבור את `verify_playbook_docx.py` (בדיקת הקבלה של
`word-export-spec.md`), ו-`tests/test_build_playbook_docx.py` בונה אותו מכל
מודל לדוגמה ומריץ עליו את הבדיקה. הבונה והבדיקה נכתבו בזמנים שונים, ובלי
הבדיקה הזו כל Word שנבנה נכשל בבדיקה של עצמו (גרסה 1.20.0).
"""
import argparse
import datetime
import io
import json
import re

import areo_lib  # noqa: F401  מוסיף את areolegal/lib לנתיב הייבוא
import basis_line
from common import rtl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from playbook import market_counts as MC, status as S, surfaces_render as R

# ---------- פלטה: נגזרת מסביבת העבודה, מותאמת לנייר ----------
INK = "16191F"
INK_SOFT = "5B6472"
LINE = "D7DBE0"
BRAND = "0B6B4F"
NAVY = "12314F"
GREEN, GREEN_BG = "0F6B4C", "E9F5F0"
AMBER, AMBER_BG = "8A6212", "FBF2DF"
RED, RED_BG = "9E2723", "FAEBEA"
SURF = "F5F7F9"

# Arial: מותקן בכל מחשב Windows ו-Mac ומכסה עברית במלואה. גופן שאינו מותקן אצל
# הנמען מוחלף בשקט, וההיררכיה של המסמך נשברת יחד איתו.
FONT = "Arial"

# שמות המסלולים לפי localization-he.md. בדיקת הקבלה מחפשת את "ירוק", "צהוב" ו"אדום" במסמך
TRACKS = {"green": ("ירוק: העמדה המועדפת", GREEN, GREEN_BG),
          "yellow": ("צהוב: הפשרה האפשרית", AMBER, AMBER_BG),
          "red": ("אדום: לא מתקבל", RED, RED_BG)}
# רמת החשיבות (importance) היא כמה הנושא חשוב לחברה, ורמת הביטחון (confidence) היא כמה
# הראיות חזקות. המפרט אוסר לגזור אחת מהשנייה, ולכן כל אחת מודפסת מהשדה שלה בלבד
IMPORTANCE = {"critical": ("קריטית", RED), "high": ("גבוהה", AMBER),
              "medium": ("בינונית", INK_SOFT), "low": ("נמוכה", INK_SOFT)}
CONFIDENCE = {"High": "גבוהה", "Medium": "בינונית", "Low": "נמוכה"}
REQ = {"Required": "חובה", "Recommended": "מומלץ", "Optional": "רשות",
       "Conditional": "מותנה", "Prohibited": "אסור"}
POLICY = {"Confirmed": "נקבע במדיניות כתובה של החברה", "Unclear": "לא ברור, טעון הכרעה"}
POSTURE = {"hold_green": "מחזיקים בירוק", "offer_yellow": "מציעים את הצהוב"}
READ_STATUS = {"read": "נקרא", "partially_read": "נקרא חלקית", "located_not_read": "אותר, לא נקרא"}
EFFECTS = [("company_effect", "ההשפעה על החברה"), ("as_discloser_effect", "כשהחברה מגלה"),
           ("as_recipient_effect", "כשהחברה מקבלת")]
# משמעות המסלול במודל היא קוד (validate_playbook_model: green=Approved וכו'), וכותרת המסלול כבר אומרת
# אותה בעברית ("ירוק: העמדה המועדפת"). קוד כזה אינו מודפס; במודל ישן meaning הוא משפט הסבר, והוא מודפס
MEANING_CODES = ("Approved", "AgreeWithComments", "NotApproved")
BRIEF_WORDS = 15
HEB = re.compile(r"[֐-׿]")
RAT = [("executive_policy_rationale", "תקציר נימוק המדיניות"),
       ("protected_interest_and_clause_purpose", "מטרת הסעיף והאינטרס המוגן"),
       ("legal_risk_and_legal_framework", "החשיפה המשפטית והמסגרת הנורמטיבית"),
       ("contractual_risk_allocation_mechanism", "מנגנון חוזי והקצאת סיכון"),
       ("company_role_and_policy_fit", "פוזיציית החברה וההתאמה למדיניות"),
       ("business_operational_impact", "ההשלכה העסקית והתפעולית"),
       ("internal_evidence_rationale", "ראיות מתוך מאגר החברה"),
       ("market_practice_assessment", "מקובלות שוק"),
       ("tradeoff_statement", "האיזון שנבחר"),
       ("related_clause_rationale", "סעיפים קשורים"),
       ("uncertainties_and_limitations", "אי-ודאות ומגבלות")]
PERSP = [("legal_analysis", "ניתוח משפטי"), ("business_analysis", "ניתוח עסקי"),
         ("internal_corpus_findings", "ממצאי מאגר פנימי"),
         ("internal_contracting_trend", "מגמת התקשרות פנימית"),
         ("external_market_trend", "מגמת שוק חיצונית")]

# ---------- כלי XML ----------
# סדר הילדים ב-w:pPr קבוע בסכמה של OOXML. אלמנט שמצורף בסוף אינו מוצג בשקט:
# כך נעלמו רקעים וקווי צבע בגרסה הראשונה, והמסמך נראה כאילו אין לו עיצוב.
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "numPr", "pBdr",
             "shd", "tabs", "bidi", "spacing", "ind", "contextualSpacing", "jc",
             "outlineLvl", "rPr"]


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def _local(tag):
    return tag.split("}")[-1]


# אותו כלל סדר חל על w:tblPr ועל w:sectPr, ושם הוא לא נאכף עד כה: bidiVisual
# שצורף בסוף tblPr נחת אחרי tblBorders ו-tblLook, ו-bidi בסוף sectPr נחת אחרי
# docGrid. Word פותח בכל זאת, אבל הקובץ אינו תקף מול הסכמה.
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
               "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
               "tblDescription", "tblPrChange"]
SECTPR_ORDER = ["footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
                "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
                "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
                "printerSettings", "sectPrChange"]


def put_ordered(parent, el, order):
    """מכניס אלמנט למקומו לפי הסכמה, ומחליף אלמנט קיים מאותו סוג."""
    name = _local(el.tag)
    old = parent.find(qn("w:" + name))
    if old is not None:
        parent.remove(old)
    idx = order.index(name) if name in order else 99
    for child in parent:
        c = _local(child.tag)
        if (order.index(c) if c in order else 99) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def ppr_put(p, el):
    """מכניס אלמנט ל-w:pPr במקומו לפי הסכמה, ומחליף אלמנט קיים מאותו סוג."""
    ppr = p._p.get_or_add_pPr()
    name = _local(el.tag)
    old = ppr.find(qn("w:" + name))
    if old is not None:
        ppr.remove(old)
    idx = PPR_ORDER.index(name) if name in PPR_ORDER else 99
    for child in ppr:
        c = _local(child.tag)
        if (PPR_ORDER.index(c) if c in PPR_ORDER else 99) > idx:
            child.addprevious(el)
            return el
    ppr.append(el)
    return el


def shade(p, fill):
    ppr_put(p, _el("w:shd", val="clear", color="auto", fill=fill))


def bar(p, color, size=18, side="start"):
    """קו צבע לצד הפסקה או מעליה."""
    ppr = p._p.get_or_add_pPr()
    b = ppr.find(qn("w:pBdr"))
    if b is None:
        b = ppr_put(p, _el("w:pBdr"))
    old = b.find(qn("w:" + side))
    if old is not None:
        b.remove(old)
    b.append(_el("w:" + side, val="single", sz=str(size), space="8", color=color))


def rtl_p(p):
    ppr_put(p, _el("w:bidi"))


def run(p, text, size=11, bold=False, color=INK, italic=False, ltr=False):
    """ריצה בעיצוב המסמך. ltr: נוסח באנגלית, כריצה משמאל לימין בתוך פסקה עברית."""
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    for k in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + k), FONT)
    if ltr:
        rtl.set_run_ltr(r)
    else:
        rpr.append(_el("w:rtl"))
    return r


def ext_link(p, text, url, size=9):
    """קישור חיצוני אמיתי (w:hyperlink עם יחס חיצוני), לכתובת המדויקת שברשומה. הכתובת אינה מודפסת."""
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    r = run(p, text, size, False, BRAND)
    r.font.underline = True
    h = _el("w:hyperlink")
    h.set(qn("r:id"), rid)
    p._p.remove(r._r)
    h.append(r._r)
    p._p.append(h)
    return h


URL_RX = re.compile(r"https?://[^\s,;)\]]+")


def record_urls(d):
    """הכתובות שברשומות המקור של המודל. קישור במסמך מוביל רק לכתובת שברשומה (בדיקת הקבלה, 7)."""
    mr = d.get("market_research") or {}
    out = set()
    for key in ("sources", "external_sources", "findings", "source_refs"):
        for x in mr.get(key) or ():
            if isinstance(x, dict):
                out |= {u.rstrip("/") for u in (x.get("url"), x.get("final_url")) if u}
    for x in list((d.get("regulatory_research") or {}).get("sources") or ()) + list(d.get("external_sources") or ()):
        if isinstance(x, dict) and x.get("url"):
            out.add(x["url"].rstrip("/"))
    return out


def linked(p, text, size=9.5, color=INK_SOFT, known=None):
    """טקסט חופשי שיש בו כתובת: כתובת שברשומה הופכת לקישור חיצוני עם שם האתר, וכתובת שאינה ברשומה
    (מקור שאותר ולא נקרא ולא נרשם) אינה מודפסת. כתובת כטקסט רגיל "אינה קישור" (word-export-spec.md)."""
    pos = 0
    for m in URL_RX.finditer(text):
        if m.start() > pos:
            run(p, text[pos:m.start()].rstrip(", "), size, False, color)
        url = m.group(0).rstrip(".")
        if known is None or url.rstrip("/") in known:
            ext_link(p, " " + re.sub(r"^https?://(www\.)?", "", url).split("/")[0], url, size)
        pos = m.start() + len(url)
    if pos < len(text):
        run(p, text[pos:], size, False, color)


def bookmark(p, name):
    """יעד לקישור פנימי: מהטבלה המקוצרת לכרטיס, ומהכרטיס לנימוק בחלק ב."""
    i = str(len(list(p.part.element.iter(qn("w:bookmarkStart")))) + 1)
    pos = 1 if p._p.pPr is not None else 0
    p._p.insert(pos, _el("w:bookmarkStart", id=i, name=name))
    p._p.append(_el("w:bookmarkEnd", id=i))


def link(p, text, anchor, size=9.5, color=BRAND):
    """קישור פנימי לחיץ אל bookmark באותו מסמך."""
    r = run(p, text, size, False, color)
    r.font.underline = True
    h = _el("w:hyperlink", anchor=anchor, history="1")
    p._p.remove(r._r)
    h.append(r._r)
    p._p.append(h)
    return h


def para(doc, text="", size=11, bold=False, color=INK, before=0, after=4, line=1.15,
         container=None, italic=False):
    p = (container or doc).add_paragraph()
    rtl_p(p)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        run(p, text, size, bold, color, italic)
    return p


def table_borders(tbl, color=LINE):
    tpr = tbl._tbl.tblPr
    b = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el("w:" + edge, val="single", sz="4", space="0", color=color))
    put_ordered(tpr, b, TBLPR_ORDER)
    put_ordered(tpr, _el("w:bidiVisual"), TBLPR_ORDER)


def cell_text(cell, text, bold=False, size=10, color=INK, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    rtl_p(p)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run(p, text, size, bold, color)
    if fill:
        cell._tc.get_or_add_tcPr().append(_el("w:shd", val="clear", color="auto", fill=fill))


# ---------- ניקוי נתונים לתצוגה ----------
def clause(v):
    """(נוסח, שפה, תורגם?) מתוך שדה נוסח במודל.

    נוסח מוצע במודל הוא {"he", "en", "source_language"}, ונוסח נסיגה {"text"}.
    מדפיסים את המחרוזת בשפת המסמך, ולעולם לא את המבנה: `str()` על המילון
    הדפיס `{'he': ...}` לתוך כל מסלול בכל כרטיס. כשיש רק אנגלית, היא
    מודפסת כריצה משמאל לימין. נוסח עברי שמקורו באנגלית מסומן "תרגום, לבדיקה".
    """
    if isinstance(v, dict):
        src = v.get("source_language")
        for k in ("he", "text", "en"):
            t = v.get(k)
            if isinstance(t, str) and t.strip():
                lang = "en" if k == "en" or not HEB.search(t) else "he"
                return t, lang, lang == "he" and src not in (None, "", "he")
        return "", "he", False
    t = txt(v)
    return t, ("he" if not t or HEB.search(t) else "en"), False


def txt(v):
    """מחרוזת מכל ערך במודל, בלי להדפיס מבנה נתונים."""
    if v is None or v is False:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, dict):
        return clause(v)[0]
    if isinstance(v, (list, tuple)):
        return "; ".join(t for t in (txt(x) for x in v) if t)
    return str(v)


def s(v, limit=None):
    """טקסט לתצוגה: רווחים מנורמלים, ובלי מקפים ארוכים, שאסורים בכל תוצר."""
    t = " ".join(txt(v).split())
    t = t.replace("—", ", ").replace("–", "-").replace(" ,", ",")
    return t[:limit] if limit else t


def words(v, n=BRIEF_WORDS):
    """שורה קצרה לטבלת "העמדות בקצרה": עד n מילים."""
    w = s(v).split()
    return " ".join(w[:n]) + ("..." if len(w) > n else "")


def roles_label(sc):
    """הפוזיציה כשהמודל נושא את התפקידים כרשימה (הסכם הדדי: "המגלה / המקבלת")."""
    return " / ".join(x["label"] for x in sc.get("company_roles") or []
                      if isinstance(x, dict) and x.get("label"))


def date_range(v):
    if isinstance(v, dict):
        a, b = s(v.get("from") or v.get("start")), s(v.get("to") or v.get("end"))
        return "%s עד %s" % (a, b) if a and b else (a or b)
    return s(v, 120)


def clean_company(v):
    """שם החברה לכריכה. פרט שטרם הושלם בפרופיל אינו מוצג על השער."""
    t = " ".join(str(v or "").split())
    for sep in (", ח.פ", " ח.פ"):
        if sep in t and "[" in t.split(sep, 1)[1][:20]:
            t = t.split(sep, 1)[0]
    return t.strip(" ,")


def short_role(v):
    """הפוזיציה בכמה מילים. במודל השדה נושא לעיתים פסקת הסבר שלמה."""
    t = " ".join(str(v or "").split())
    for sep in (" — ", " - ", "(", ",", ";", "."):
        if sep in t:
            t = t.split(sep)[0]
    return t.strip() or "לא צוין"


def produced_on(d):
    """תאריך ההפקה בפורמט יום.חודש.שנה. כשחסר במודל, תאריך היום."""
    raw = (d.get("generated_at") or d.get("last_updated") or "")[:10]
    if len(raw) == 10 and raw[4] == "-":
        y, mo, day = raw.split("-")
        return "%s.%s.%s" % (day, mo, y)
    return datetime.date.today().strftime("%d.%m.%Y")


# ---------- בלוקי המסמך ----------
def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:cs"), FONT)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(1.8)
        put_ordered(sec._sectPr, _el("w:bidi"), SECTPR_ORDER)
        put_ordered(sec._sectPr, _el("w:rtlGutter"), SECTPR_ORDER)


SOURCES_HEAD = "המקורות שעליהם נשען הפלייבוק"
NOT_READ = ("located_not_read", "excluded", "blocked")
DRAFTED_HE = {"same_side": "הצד שלנו", "other_side": "הצד שכנגד", "negotiated": "נוהל עליו משא ומתן",
              "neutral_model": "מודל מאוזן", "unknown": "לא ידוע"}
POINT_HE = (("ours", "לטובתנו"), ("middle", "באמצע"), ("theirs", "לטובתם"))


def he_date(v):
    raw = str(v or "")[:10]
    if len(raw) == 10 and raw[4] == "-":
        y, mo, day = raw.split("-")
        return "%s.%s.%s" % (day, mo, y)
    return s(v, 20)


def table(doc, heads, rows, widths):
    """טבלה מימין לשמאל. תא שהוא (טקסט, כתובת) הוא קישור חיצוני."""
    t = doc.add_table(rows=1, cols=len(heads))
    table_borders(t)
    for i, head in enumerate(heads):
        cell_text(t.rows[0].cells[i], head, True, 9, NAVY, SURF)
    for row in rows:
        c = t.add_row().cells
        for i, v in enumerate(row):
            if isinstance(v, tuple):
                c[i].text = ""
                p = c[i].paragraphs[0]
                rtl_p(p)
                ext_link(p, v[0], v[1], 8.5)
            else:
                cell_text(c[i], v, False, 8.5)
    for r in t.rows:
        for cell, w in zip(r.cells, widths):
            cell.width = Cm(w)
    para(doc, after=8)


def sources_chapter(doc, d, num):
    """"המקורות שעליהם נשען הפלייבוק" (word-export-spec.md, "סדר המסמך"), כשרץ מחקר שוק או מחקר דין:
    המשפט הקבוע, טבלת המקורות עם קישור לחיץ, הטווח לכל נושא, מה לא אומת, מקורות הדין והמגבלות.
    בלי הפרק הזה Word של פלייבוק שנשען על השוק נכשל בבדיקת הקבלה, ולא נמסר (הרצה חיה, 20.10.2026)."""
    mr = d.get("market_research") or {}
    rr = d.get("regulatory_research") or {}
    refs = [r for r in mr.get("source_refs") or () if isinstance(r, dict) and r.get("read_status") not in NOT_READ]
    legal = [x for x in rr.get("sources") or () if isinstance(x, dict)]
    if not refs and not legal:
        return False
    h1(doc, SOURCES_HEAD, num)
    # המשפט נגזר מהספירה בכל הפקה (market_research.fixed_sentence), ולא מועתק מהמודל: משפט שנשמר במודל
    # יכול להיות ישן, וזה בדיוק "נבנה לפי 10 הסכמים, מ-10 גופים" כשנספרו 6 (סבב 3, ממצא 9)
    import market_research
    sentence = market_research.fixed_sentence(d) or MC.basis_sentence(d)
    if sentence:
        para(doc, s(sentence), 10.5, after=4)
    if mr.get("research_date"):
        para(doc, "תאריך המחקר: %s" % he_date(mr["research_date"]), 9.5, color=INK_SOFT, after=8)
    if refs:
        h3(doc, "מקורות השוק")
        table(doc, ["#", "גוף", "מסמך", "תאריך", "ניסח / לטובת מי", "קישור"],
              [[s(r.get("source_id"), 8), s(r.get("issuing_body"), 60), s(r.get("document"), 80),
                he_date(r.get("doc_date")) or "לא צוין",
                DRAFTED_HE.get(r.get("drafted_by"), "לא ידוע"),
                ("פתיחה", r.get("final_url") or r.get("url")) if (r.get("final_url") or r.get("url")) else ""]
               for r in refs], (1.0, 3.2, 4.6, 1.8, 2.8, 1.6))
    ranges = [x for x in mr.get("ranges") or () if isinstance(x, dict)]
    if ranges:
        h3(doc, "הטווח לכל נושא")
        for rng in ranges:
            p = para(doc, after=1)
            run(p, s(rng.get("topic") or rng.get("rule_id")), 10, True, INK)
            for key, label in POINT_HE:
                pt = rng.get(key)
                if isinstance(pt, dict) and pt.get("source_ids"):
                    text = "%s, מקורות: %s" % (s(pt.get("certainty")) or "ודאות לא צוינה",
                                               ", ".join(pt["source_ids"]))
                elif key == "theirs" and rng.get("theirs_note"):
                    text = s(rng["theirs_note"])
                else:
                    text = "ממצא לא מאומת, פחות משלושה מקורות"
                q = para(doc, after=1)
                run(q, label + ": ", 9.5, True, INK_SOFT)
                run(q, text, 9.5, False, INK)
    open_items = [s(u.get("text") or u.get("label")) for u in mr.get("unverified_findings") or ()
                  if isinstance(u, dict)]
    open_items += ["%s: טעון אימות (%s)" % (s(x.get("title")), "אותר, לא נקרא" if x.get("read_status") ==
                   "located_not_read" else "נקרא חלקית")
                   for x in legal if x.get("read_status") in ("located_not_read", "partially_read")]
    if open_items:
        h3(doc, "מה לא אומת")
        for x in open_items:
            q = para(doc, after=3)
            bar(q, AMBER, 12)
            linked(q, x, known=record_urls(d))
    if legal:
        h3(doc, "מקורות הדין")
        table(doc, ["#", "מקור", "גוף", "נקרא", "קישור"],
              [[s(x.get("source_id"), 8), s(x.get("title"), 90), s(x.get("issuer"), 40),
                {"read": "נקרא", "partially_read": "נקרא חלקית"}.get(x.get("read_status"), "אותר, לא נקרא"),
                ("פתיחה", x["url"]) if x.get("url") else ""]
               for x in legal], (1.0, 6.4, 3.4, 2.4, 1.6))
    lim = mr.get("limitations")
    if lim:
        h3(doc, "מגבלות")
        for x in ([lim] if isinstance(lim, str) else lim):
            linked(para(doc, after=3), s(x), known=record_urls(d))
    return True


def h1(doc, text, num=""):
    p = para(doc, before=20, after=2)
    if num:
        run(p, num + "   ", 18, True, BRAND)
    run(p, text, 18, True, NAVY)
    bar(para(doc, after=10), LINE, 8, "top")


def h2(doc, text, anchor=None):
    p = para(doc, before=14, after=4)
    run(p, text, 13.5, True, NAVY)
    if anchor:
        bookmark(p, anchor)
    return p


def status_para(container, tag, text, size, before=0, after=4):
    """טקסט מעמד בתוך בקר תוכן עם תגית (word-export-spec.md, "מעמד שמתעדכן בלי
    לבנות את המסמך מחדש"): אחרי קליטת גיליון העמדות `restamp_docx` מעדכן אותו
    במקום, בלי לבנות את המסמך מחדש. הפסקה מכילה את הטקסט לבדו, כי כך הבדיקה
    מזהה תווית."""
    p = container.add_paragraph()
    rtl_p(p)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    R.add_status_control(p, tag, text, size)
    return p


def h3(doc, text):
    run(para(doc, before=9, after=2), text, 11, True, INK_SOFT)


def fam_role(fams):
    for f in fams:
        if isinstance(f, dict) and f.get("company_role"):
            return f["company_role"]
    return ""


def corpus_label(model, scope):
    """מה שמוצג בטבלת תמונת המצב: הסכמי החברה, ולצדם מספר הקבצים כשהוא שונה.

    "של החברה" צמוד למספר, כדי שלא ייקרא כמספר ההסכמים ממחקר השוק: "0 מסמכים"
    בפלייבוק שנבנה מ-8 הסכמים שפורסמו ברשת הוא סתירה, ובדיקת הקבלה עוצרת עליו.
    """
    units = basis_line.agreement_units(model)
    files = (scope.get("corpus_basis") or {}).get("files")
    if units is None:
        units = scope.get("total_document_count")
        if not isinstance(units, int):
            return "לא צוין"
    if units == 0:
        return "אין הסכמים של החברה מהסוג הזה"
    if isinstance(files, int) and files != units:
        return "%d הסכמים של החברה (מתוך %d קבצים)" % (units, files)
    return "%d הסכמים של החברה" % units


def cover(doc, m):
    run(para(doc, before=150, after=2), m["company"], 13, True, BRAND)
    run(para(doc, after=0), "פלייבוק חוזי", 30, True, NAVY)
    run(para(doc, after=6), m["family"], 16, False, INK_SOFT)
    # התווית (status_label) על השער, מילה במילה, כמו בכותרת העליונה בכל עמוד
    status_para(doc, R.STATUS_TAG, m["st"]["label"], 11, after=14)
    bar(para(doc, after=16), BRAND, 26, "top")
    run(para(doc, after=6), {"one": "חלק א: המדיניות. חלק ב: הנימוקים והנספחים", True: "הגרסה המלאה",
                             False: "הגרסה המקוצרת"}[m["full"]], 11, True, BRAND)
    # שורת הבסיס: על כמה הסכמים הפלייבוק נשען, ובאיזה מסלול נבנה. אותו
    # נוסח בדיוק בפס הבסיס בארטיפקט, ולכן הוא נגזר ב-`basis_line.py`
    # ואינו מנוסח כאן. הוא אינו מספר הקבצים בתיקייה
    if m["basis"]:
        run(para(doc, after=4), m["basis"], 10.5, False, INK)
    if m["mostly_drafts"]:
        run(para(doc, after=4), "פלייבוק ראשוני, נבנה ברובו מטיוטות", 10.5, True, INK_SOFT)
    rows = [("פוזיציית החברה", m["role"]), ("גרסה", m["version"]),
            ("הופק בתאריך", m["date"])]
    if m["updated"] and m["updated"] != m["date"]:
        rows.append(("עדכון אחרון", m["updated"]))
    for k, v in rows:
        p = para(doc, after=3)
        run(p, k + ": ", 10.5, True, INK_SOFT)
        run(p, v, 10.5, False, INK)
    p = para(doc, before=18, after=0)
    shade(p, SURF)
    bar(p, BRAND, 18)
    run(p, "חסוי. הוכן לצורך קבלת ייעוץ משפטי ומיועד לשימוש פנימי בלבד.", 9.5, True, INK_SOFT)
    p = para(doc, after=0)
    shade(p, SURF)
    run(p, "טיוטה מקצועית. חובה בדיקת עורך דין לפני הסתמכות או שליחה.", 9.5, False, INK_SOFT)
    p.runs[-1].add_break(WD_BREAK.PAGE)


def kv(doc, rows):
    """טבלת שתי עמודות. ערך שהוא (תגית, טקסט) נכתב בבקר תוכן, כמו כל טקסט מעמד."""
    t = doc.add_table(rows=0, cols=2)
    table_borders(t)
    for k, v in rows:
        c = t.add_row().cells
        cell_text(c[0], k, True, 10, INK_SOFT, SURF)
        if isinstance(v, tuple):
            c[1].text = ""
            p = c[1].paragraphs[0]
            rtl_p(p)
            R.add_status_control(p, v[0], v[1], 10)
        else:
            cell_text(c[1], v, False, 10)
        c[0].width = Cm(4.6)
        c[1].width = Cm(11.2)
    para(doc, after=8)


def ev_table(doc, ev):
    t = doc.add_table(rows=1, cols=3)
    table_borders(t)
    for i, head in enumerate(["ההסכם", "סעיף", "הנוסח שנמצא"]):
        cell_text(t.rows[0].cells[i], head, True, 9.5, NAVY, SURF)
    for e in ev:
        c = t.add_row().cells
        cell_text(c[0], s(e.get("agreement") or e.get("source_document"), 40), False, 9)
        cell_text(c[1], s(e.get("clause_ref"), 18), False, 9)
        cell_text(c[2], s(e.get("quote") or e.get("text")), False, 9, INK_SOFT)
    for row in t.rows:
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(2.0)
        row.cells[2].width = Cm(9.8)
    para(doc, after=8)


def anchors(doc, r, limit):
    srcs = [s(x, limit) for x in (r.get("regulatory_sources") or []) if x]
    if not srcs:
        return
    h3(doc, "עיגונים בדין")
    for a in srcs:
        p = para(doc, after=2)
        bar(p, LINE, 8)
        run(p, a, 9.5, False, INK_SOFT)


def legal_review(r):
    """האם הכלל טעון בדיקה משפטית: רשימת סעיפים במודל החדש, או דגל במודל ישן."""
    v = r.get("legal_review_required")
    if isinstance(v, list):
        return any(not (isinstance(x, dict) and x.get("checked_by")) for x in v)
    return bool(v)


def status_line(doc, r, st):
    """שורת המצב של הכרטיס: חשיבות, כלל נוכחות, סוג הסעיף, מעמד המדיניות, רמת
    הביטחון, ומזהה הכלל. שדה שאין לו ערך במודל אינו מודפס."""
    p = para(doc, after=4)
    shade(p, SURF)
    first = [True]

    def item(label, value, color=INK, bold=False):
        if not value:
            return
        run(p, ("" if first[0] else "     ") + label + "  ", 9.5, True, INK_SOFT)
        run(p, value, 9.5, bold, color)
        first[0] = False

    imp = IMPORTANCE.get(str(r.get("importance") or "").lower())
    if imp:
        item("רמת חשיבות", imp[0], imp[1], True)
    pres = r.get("presence_rule") or {}
    req = pres.get("requirement") if isinstance(pres, dict) else pres
    if req:
        item("כלל נוכחות", REQ.get(req, s(req, 30)))
    item("סוג הסעיף", s(r.get("clause_kind"), 20))
    if r.get("policy_status") == "Confirmed":
        item("מעמד", POLICY["Confirmed"])
    elif r.get("policy_status") == "Unclear":
        item("מעמד", POLICY["Unclear"])
    else:
        # "מאושר" אינו מודפס ליד כלל לפני שהגיליון החתום נקלט (המפרט, "שורת מצב")
        item("מעמד", "מוצע, טעון אישור")
    conf = CONFIDENCE.get(s(r.get("confidence") or r.get("confidence_level")))
    item("רמת ביטחון", conf)
    if legal_review(r):
        item("", "טעון בדיקה משפטית", RED, True)
    if r.get("rule_id"):
        item("מזהה", s(r["rule_id"], 20), INK_SOFT)
    posture = POSTURE.get(r.get("round1_posture"))
    if posture:
        q = para(doc, after=8)
        run(q, "בסבב הראשון: ", 9.5, True, INK_SOFT)
        run(q, posture, 9.5, False, INK)


def track(doc, r, key, st):
    """מסלול אחד: הכותרת, הנוסח המוצע במלואו, העמדה, ההשפעה על החברה, שורת המקור
    ומעמד העמדה. מחזיר False כשאין במסלול שום נתון."""
    blk = (r.get("traffic_light") or {}).get(key) or {}
    text, lang, translated = clause(blk.get("proposed_clause") or blk.get("fallback_wording"))
    meaning = "" if blk.get("meaning") in MEANING_CODES else s(blk.get("meaning"))
    position = s(blk.get("position"))
    if not (text or meaning or position):
        return False
    title, color, fill = TRACKS[key]

    def line(after=2):
        q = para(doc, after=after)
        bar(q, color, 22)
        shade(q, fill)
        return q

    p = para(doc, before=8, after=2)
    bar(p, color, 22)
    shade(p, fill)
    run(p, title, 11, True, color)
    # מתחת לכותרת הצהוב: מאיפה נלמד קו הנסיגה (המפרט, "הצהוב תמיד קיים")
    origin = (blk.get("fallback_origin") or r.get("fallback_origin") or {}) if key == "yellow" else {}
    if isinstance(origin, dict) and origin.get("label_he"):
        run(line(), s(origin["label_he"]), 9.5, False, INK_SOFT, italic=True)
    # הנוסח המוצע לעולם אינו מקוצר: זה הטקסט שעורך הדין מעתיק אל ההסכם
    body = " ".join(text.split()).replace("—", ", ").replace("–", "-") if text else meaning
    if body:
        run(line(), body, 10.5, False, INK, ltr=(lang == "en" and bool(text)))
        if translated:
            run(line(), "תרגום, לבדיקה", 9, True, INK_SOFT)
    if position and position != body:
        q = line()
        run(q, "העמדה: ", 9.5, True, INK_SOFT)
        run(q, position, 9.5, False, INK)
    if meaning and text:
        run(line(), meaning, 9.5, False, INK_SOFT, italic=True)
    for field, label in EFFECTS:
        if blk.get(field):
            q = line()
            run(q, label + ": ", 9.5, True, INK_SOFT)
            run(q, s(blk[field]), 9.5, False, INK)
    for v in (blk.get("role_effects") or {}).values() if isinstance(blk.get("role_effects"), dict) else ():
        if s(v):
            q = line()
            run(q, "ההשפעה על החברה: ", 9.5, True, INK_SOFT)
            run(q, s(v), 9.5, False, INK)
    src = s(blk.get("source_label") or blk.get("label"))
    if src:
        q = line()
        run(q, "מקור: ", 9, True, INK_SOFT)
        run(q, src, 9, False, INK_SOFT)
    if r.get("rule_id"):
        k = S.cell_key(r["rule_id"], key)
        q = line(after=6)
        R.add_status_control(q, R.CELL_TAG + k, R.cell_status_text(None, st, k), 8.5)
    return True


def has_rationale(r):
    pr = r.get("policy_rationale") or {}
    cp = r.get("company_perspective") or {}
    return bool(any(pr.get(k) for k, _ in RAT) or any(cp.get(k) for k, _ in PERSP)
                or r.get("regulatory_sources") or r.get("evidence"))


def topic(doc, r, n, full, st):
    h2(doc, "%d. %s" % (n, s(r.get("topic_name"))), "card_%d" % n)
    status_line(doc, r, st)
    if r.get("clause_issue"):
        para(doc, s(r["clause_issue"]), 10.5, after=8, color=INK_SOFT)
    for key in ("green", "yellow", "red"):
        track(doc, r, key, st)

    pr = r.get("policy_rationale") or {}
    if full == "card":
        if pr.get("executive_policy_rationale"):
            h3(doc, "נימוק משפטי")
            para(doc, s(pr["executive_policy_rationale"]), 10.5, after=6)
        anchors(doc, r, 160)
        if has_rationale(r):
            p = para(doc, after=6)
            link(p, "הנימוק המלא והראיות: בחלק ב, נושא %d" % n, "why_%d" % n)
    elif not full:
        if pr.get("executive_policy_rationale"):
            h3(doc, "נימוק משפטי")
            para(doc, s(pr["executive_policy_rationale"]), 10.5, after=6)
        anchors(doc, r, 160)
    else:
        rationale(doc, r, pr)
    bar(para(doc, before=6, after=10), LINE, 6, "top")


def brief(doc, rules, num):
    """"העמדות בקצרה" (word-export-spec.md, "סדר המסמך"): שורה לכל כלל, הירוק, הצהוב
    והאדום בשורה קצרה מתוך `position`, איך פותחים בסבב הראשון, וקישור לכרטיס."""
    h1(doc, "העמדות בקצרה", num)
    t = doc.add_table(rows=1, cols=5)
    table_borders(t)
    for i, head in enumerate(["הנושא", "ירוק", "צהוב", "אדום", "בסבב הראשון"]):
        cell_text(t.rows[0].cells[i], head, True, 9.5, NAVY, SURF)
    for n, r in enumerate(rules, 1):
        c = t.add_row().cells
        c[0].text = ""
        p = c[0].paragraphs[0]
        rtl_p(p)
        link(p, "%d. %s" % (n, s(r.get("topic_name"))), "card_%d" % n, 9)
        if r.get("rule_id"):
            run(p, " (%s)" % s(r["rule_id"], 20), 8, False, INK_SOFT)
        tl = r.get("traffic_light") or {}
        for j, key in enumerate(("green", "yellow", "red"), 1):
            blk = tl.get(key) or {}
            fallback = "" if blk.get("meaning") in MEANING_CODES else blk.get("meaning")
            cell_text(c[j], words(blk.get("position") or fallback), False, 9)
        cell_text(c[4], POSTURE.get(r.get("round1_posture"), ""), False, 9)
    for row in t.rows:
        for cell, w in zip(row.cells, (4.0, 3.4, 3.4, 3.4, 1.9)):
            cell.width = Cm(w)
    para(doc, after=8)


def rationale(doc, r, pr):
    """הנימוק המלא של נושא: פרספקטיבת החברה, מדוע נבחרה המדיניות, העיגונים והראיות."""
    cp = r.get("company_perspective") or {}
    if any(cp.get(k) for k, _ in PERSP):
        h3(doc, "פרספקטיבת החברה")
        for k, label in PERSP:
            if cp.get(k):
                p = para(doc, after=4)
                run(p, label + ": ", 10, True, INK_SOFT)
                run(p, s(cp[k]), 10, False, INK)
    if any(pr.get(k) for k, _ in RAT):
        h3(doc, "מדוע נבחרה מדיניות זו")
        for k, label in RAT:
            if pr.get(k):
                p = para(doc, after=4)
                run(p, label + ": ", 10, True, INK_SOFT)
                run(p, s(pr[k]), 10, False, INK)
    anchors(doc, r, None)
    if r.get("evidence"):
        h3(doc, "הראיות מהמאגר")
        ev_table(doc, r["evidence"])


def header_footer(doc, m):
    """בכותרת העליונה של כל עמוד: שם הפלייבוק והתווית (המפרט, "בכל עמוד").
    בכותרת התחתונה: חסוי ומספר עמוד."""
    for sec in doc.sections:
        p = sec.header.paragraphs[0]
        rtl_p(p)
        run(p, "%s · פלייבוק %s" % (m["company"], m["family"]), 8.5, True, INK_SOFT)
        status_para(sec.header, R.STATUS_TAG, m["st"]["label"], 8, after=0)
        p = sec.footer.paragraphs[0]
        rtl_p(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, "%s · %s · חסוי · עמוד " % (m["company"], m["family"]), 8.5, False, INK_SOFT)
        p._p.append(_el("w:fldSimple", instr="PAGE"))


def build(d, prof, out, full):
    sc = d.get("scope_summary") or {}
    fams = sc.get("contract_families") or []
    st = S.derive_status(d)
    role = sc.get("company_role") or fam_role(fams) or roles_label(sc)
    m = {"company": clean_company(prof.get("organization_name")) or "החברה",
         "family": ", ".join(f.get("contract_family_name", "") if isinstance(f, dict) else str(f)
                             for f in fams) or "הסכמי החברה",
         # הפוזיציה חובה על השער. במודל היא יושבת על סוג ההסכם, ולכן
         # כשאינה בראש התחולה היא נלקחת משם
         "role": short_role(role),
         "role_full": s(role, 600),
         "version": str(d.get("version") or "1.0"),
         "date": produced_on(d),
         "corpus": corpus_label(d, sc),
         "basis": basis_line.line_for(d) or "",
         "mostly_drafts": basis_line.mostly_drafts(d),
         "updated": produced_on({"last_updated": d.get("last_updated")}) if d.get("last_updated") else "",
         "full": full, "st": st}
    doc = Document()
    base(doc)
    cover(doc, m)
    header_footer(doc, m)

    h1(doc, "תמונת מצב", "1")
    rows = [("סוג ההסכמים", m["family"]),
            ("פוזיציית החברה", m["role_full"] or m["role"]),
            ("נושאי מדיניות", str(len(d["rules"]))),
            ("הסכמי החברה שנותחו", m["corpus"])]
    # מחקר השוק: המשפט נגזר מהספירה (market_counts), ולא מנוסח כאן
    if MC.basis_sentence(d):
        rows.append(("מחקר שוק", MC.basis_sentence(d)))
    rows += [("טווח תאריכים", date_range(sc.get("date_range")) or "לא צוין"),
             # המעמד נגזר בקוד (playbook/status.py), ואינו הקוד באנגלית שבמודל ("Draft")
             ("מעמד הפלייבוק", (R.STATUS_TAG, st["label"]))]
    kv(doc, rows)
    status_para(doc, R.SUMMARY_TAG, R.summary_text(st), 10, after=8)
    lim = sc.get("limitations")
    if lim:
        h3(doc, "מגבלות")
        for x in ([lim] if isinstance(lim, str) else lim):
            q = para(doc, after=4)
            bar(q, AMBER, 12)
            run(q, s(x), 10, False, INK_SOFT)

    n = 2
    if sources_chapter(doc, d, str(n)):
        n += 1
    brief(doc, d["rules"], str(n))
    n += 1
    if full == "one":
        # קובץ אחד בשני חלקים (word-export-spec.md, "קובץ Word אחד לכל פלייבוק, בשני חלקים"):
        # חלק א הוא מה שעורך דין פותח באמצע משא ומתן, וחלק ב הנימוקים והנספחים
        h1(doc, "חלק א: המדיניות לפי נושא", str(n))
        for i, r in enumerate(d["rules"], 1):
            topic(doc, r, i, "card", st)
        h1(doc, "חלק ב: הנימוקים", str(n + 1))
        with_why = [(i, r) for i, r in enumerate(d["rules"], 1) if has_rationale(r)]
        if not with_why:
            para(doc, "אין במודל נימוק מלא מעבר למה שמודפס בכרטיסים.", 10.5, color=INK_SOFT)
        for i, r in with_why:
            h2(doc, "%d. %s" % (i, s(r.get("topic_name"))), "why_%d" % i)
            rationale(doc, r, r.get("policy_rationale") or {})
            bar(para(doc, before=6, after=10), LINE, 6, "top")
    else:
        h1(doc, "מדיניות לפי נושא", str(n))
        for i, r in enumerate(d["rules"], 1):
            topic(doc, r, i, full, st)

    if full and d.get("source_log"):
        h1(doc, "נספח: ההסכמים שנותחו", str(n + 2 if full == "one" else n + 1))
        t = doc.add_table(rows=1, cols=3)
        table_borders(t)
        for i, head in enumerate(["ההסכם", "הצד שכנגד", "מעמד הקריאה"]):
            cell_text(t.rows[0].cells[i], head, True, 9.5, NAVY, SURF)
        for src in d["source_log"]:
            if not isinstance(src, dict):
                continue
            c = t.add_row().cells
            cell_text(c[0], s(src.get("agreement_name") or src.get("document_name")
                              or src.get("title") or src.get("file"), 40), False, 9)
            cell_text(c[1], s(src.get("counterparty"), 40) or "לא צוין", False, 9)
            rs = src.get("read_status")
            cell_text(c[2], READ_STATUS.get(rs) or s(rs, 40) or "לא צוין", False, 9)
    # רשת הביטחון של הספרייה: כיווניות בכל פסקה, ריצה, טבלה ומקטע, ומאפייני קובץ
    # נקיים (בלי "python-docx" ובלי תאריכי 2013 של התבנית). המחבר הוא שם החברה
    author = m["company"] if m["company"] != "החברה" else ""
    rtl.enforce_document(doc, author=author, title="פלייבוק חוזי - %s" % m["family"])
    doc.save(out)


def main():
    ap = argparse.ArgumentParser(description="בונה את מסמך הפלייבוק ב-Word מהמודל הקנוני")
    ap.add_argument("--model", required=True, help="playbook-model.json")
    ap.add_argument("--profile", required=True, help="פרופיל-לקוח.json")
    ap.add_argument("--version", default="one", choices=["one", "short", "full"],
                    help="one (ברירת המחדל) = קובץ אחד בשני חלקים: המדיניות, ואחריה הנימוקים והנספחים. "
                         "short ו-full נשארו לפלייבוקים ישנים")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    d = json.load(io.open(a.model, encoding="utf-8"))
    prof = json.load(io.open(a.profile, encoding="utf-8"))
    if not d.get("rules"):
        print("במודל אין כללים. אין ממה לבנות מסמך.")
        return 1
    build(d, prof, a.out, {"one": "one", "full": True, "short": False}[a.version])
    print("נבנה %s (%s, %d נושאים)" % (a.out, {"one": "קובץ אחד בשני חלקים", "full": "מלאה",
                                               "short": "מקוצרת"}[a.version], len(d["rules"])))
    print("לפני מסירה: verify_playbook_docx.py על הקובץ. קוד יציאה 3 = אין למסור.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

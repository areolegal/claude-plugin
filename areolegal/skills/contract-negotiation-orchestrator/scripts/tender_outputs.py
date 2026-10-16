#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""התוצרים במכרז ובתנאים מקוונים: טבלת שאלות הבהרה ותזכיר החלטה פנימי, בקובצי Word.

במכרז לפני הגשה ובתנאים מקוונים אין סימון בגוף החוזה ואין בועות: שינוי בגוף
עלול לפסול את ההצעה, והערוץ היחיד הוא שאלות ההבהרה (`overview.md`, שלב ד;
הערות הנושא ונותן, שינוי 7). `apply_redlines.py` מסרב לרוץ על נייר כזה.

* **`tender_pre_bid`**: שני קבצים. **טבלת שאלות הבהרה** (מספר, מסמך, סעיף,
  עמוד, השאלה), ממוינת לפי חשיבות ועד 20 שאלות, מ-`clarification_questions`
  שבקובץ המצב. כל שאלה מנוסחת כבקשה; אין בה "מוסכם", "לא מוסכם", מדיניות,
  פלייבוק או נימוק פנימי, כי התשובות מתפרסמות לכל המציעים, ו-`internal_reason`
  לעולם אינו נכנס לטבלה. ו**תזכיר החלטה פנימי (go/no-go)**, שאינו נשלח, ובראשו
  המועד האחרון לשאלות ההבהרה והמועד האחרון להגשה.
* **`tender_awarded`** ו-**`standard_form_online`**: התזכיר הפנימי בלבד.

    python3 scripts/tender_outputs.py --deal "<תיקיית העסקה>" --memo <תזכיר.json> \\
        --agreement-type "<סוג ההסכם>" --company "<שם החברה>"

`--memo` הוא JSON: `{"questions_deadline", "submission_deadline", "items": [{"topic",
"contract_ref", "what", "nature": "mandatory_law | commercial", "price_impact", "risk_impact",
"recommendation": "go | go_priced | go_if_answered | no_go", "note"}]}`.

קוד יציאה: **0** נבנה, **3** שאלה שאסור שתצא לכל המציעים (מסורבת, ולא נכתב
דבר), **2** קלט חסר או סוג נייר שאינו מכרז.
"""
import argparse
import datetime as _dt
import io
import json
import os
import re
import sys

import areo_lib  # noqa: F401  מוסיף את areolegal/lib לנתיב הייבוא

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from common import paths, rtl  # noqa: E402

MAX_QUESTIONS = 20
PRIORITY = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
FORBIDDEN = re.compile(r"מוסכם|לא מוסכם|מדיניות|פלייבוק|playbook|policy|קו אדום|red line|נסיגה|fallback|"
                       r"\bagreed\b|not agreed|עמדת החברה|our position", re.I)
RECOMMEND = {"go": "לגשת", "go_priced": "לגשת ולתמחר את הסיכון", "go_if_answered": "לגשת בכפוף לתשובה לשאלה",
             "no_go": "לא לגשת"}
NATURE = {"mandatory_law": "הוראת דין כופה", "commercial": "עמדה מסחרית של המזמין"}


def load(path, what):
    if not path or not os.path.exists(path):
        print("%s אינו קיים: %s" % (what, path))
        raise SystemExit(2)
    with io.open(path, encoding="utf-8") as fh:
        return json.load(fh)


def table(doc, head, rows):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = h
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v or "")
    rtl.set_table_rtl(t)
    return t


def out_path(deal, agreement_type, company, what, date):
    name = " - ".join(x for x in (agreement_type, company, what, date) if x) + ".docx"
    name = re.sub(r'[\\/:*?"<>|]', "", name)
    p = paths.unique_path(os.path.join(deal, name))
    paths.check_write_path(p)
    return p


def build_questions(qs, path, company):
    doc = Document()
    rtl.add_paragraph(doc, "טבלת שאלות הבהרה", bold=True, size=16)
    rows = []
    for n, q in enumerate(qs, 1):
        rows.append((n, q.get("document"), q.get("contract_ref"), q.get("page"), q.get("question")))
    table(doc, ("מספר", "מסמך", "סעיף", "עמוד", "השאלה"), rows)
    rtl.enforce_document(doc, author=company)
    doc.save(path)


def build_memo(memo, path, company, paper):
    doc = Document()
    rtl.add_paragraph(doc, "תזכיר החלטה פנימי: האם לגשת", bold=True, size=16)
    rtl.add_paragraph(doc, "פנימי, לא למשלוח", bold=True, size=11)
    if memo.get("questions_deadline"):
        rtl.add_paragraph(doc, "המועד האחרון לשאלות הבהרה: %s" % memo["questions_deadline"], bold=True)
    if memo.get("submission_deadline"):
        rtl.add_paragraph(doc, "המועד האחרון להגשה: %s" % memo["submission_deadline"], bold=True)
    if paper == "standard_form_online" and memo.get("signed_agreement_recommendation"):
        rtl.add_paragraph(doc, memo["signed_agreement_recommendation"])
    rows = []
    for it in memo.get("items") or []:
        rows.append((it.get("topic"), it.get("contract_ref"), it.get("what"), NATURE.get(it.get("nature"), ""),
                     it.get("price_impact"), it.get("risk_impact"),
                     RECOMMEND.get(it.get("recommendation"), it.get("recommendation") or "")))
    table(doc, ("נושא", "סעיף", "מה קובע הסעיף", "דין כופה או עמדה מסחרית", "השפעה על המחיר", "השפעה על הסיכון",
                "המלצה"), rows)
    for p in doc.sections[0].header.paragraphs or [doc.sections[0].header.add_paragraph()]:
        p.add_run("פנימי, לא למשלוח").font.size = Pt(9)
    rtl.enforce_document(doc, author=company)
    doc.save(path)


def main(argv=None):
    sys.dont_write_bytecode = True
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--deal", required=True)
    ap.add_argument("--memo", required=True, help="JSON התזכיר")
    ap.add_argument("--agreement-type", required=True)
    ap.add_argument("--company", default="")
    ap.add_argument("--date", help="DD.MM.YYYY; בלעדיו: היום")
    a = ap.parse_args(argv)
    state = load(os.path.join(a.deal, "negotiation_state.json"), "קובץ המצב")
    paper = state.get("paper_type") or (state.get("meta") or {}).get("paper_type")
    if paper not in ("tender_pre_bid", "tender_awarded", "standard_form_online"):
        print("סוג הנייר הוא %s, ולא מכרז או תנאים מקוונים. התגובה נבנית ב-apply_redlines.py." % paper)
        return 2
    memo = load(a.memo, "התזכיר")
    date = a.date or _dt.date.today().strftime("%d.%m.%Y")
    written = []
    if paper == "tender_pre_bid":
        qs = sorted(state.get("clarification_questions") or [],
                    key=lambda q: (PRIORITY.get(str(q.get("priority")).upper(), 3), q.get("number") or 0))
        bad = [q for q in qs if FORBIDDEN.search(q.get("question") or "")]
        if bad:
            for q in bad:
                print("שאלה %s נוסחה כעמדה ולא כבקשה (\"%s\"): התשובות מתפרסמות לכל המציעים."
                      % (q.get("number"), FORBIDDEN.search(q["question"]).group(0)))
            return 3
        if len(qs) > MAX_QUESTIONS:
            print("יש %d שאלות; בטבלה נכנסות %d המהותיות ביותר, לפי החשיבות." % (len(qs), MAX_QUESTIONS))
            qs = qs[:MAX_QUESTIONS]
        p = out_path(a.deal, a.agreement_type, a.company, "שאלות הבהרה", date)
        build_questions(qs, p, a.company)
        written.append(p)
    p = out_path(a.deal, a.agreement_type, a.company, "תזכיר פנימי", date)
    build_memo(memo, p, a.company, paper)
    written.append(p)
    for p in written:
        problems = rtl.audit_docx(p, run_date=_dt.date.today())
        print("נבנה: %s%s" % (p, "" if not problems else "  (בדיקת כיווניות: %s)" % "; ".join(problems[:3])))
    return 0


if __name__ == "__main__":
    sys.exit(main())

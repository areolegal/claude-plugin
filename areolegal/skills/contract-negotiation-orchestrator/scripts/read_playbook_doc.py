#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""קורא פלייבוק של הלקוח ממסמך Word, PDF או Excel, בלי לכתוב דבר לתיקיית הלקוח.

מסמך של הלקוח משמש כמו שהוא (`overview.md`, "מקור העמדות", "יש פלייבוק:
שימוש"; הערות הנושא ונותן, שינוי 45). הקורא מחזיר את מה שבמסמך, בלי להשלים
ובלי לפרש: הנושאים, המסלולים ותנאיהם, סימני האדום, הנוסחים החלופיים, הסגמנטים,
וכללי ההפעלה שהמסמך קובע לעצמו (פרק "הוראות הפעלה"). שדה שאין לו מקור במסמך
נשאר ריק, וההחלה ממלאת אותו לפי `decision_matrix.md`, ולא הקורא.

איך הוא קורא:
* **טבלה** שבשורת הכותרת שלה עמודות של מסלולים (ירוק, צהוב, אדום; Green,
  Yellow, Red; מועדף, נסיגה, קו אדום): כל שורה היא נושא.
* **כותרות**: כל כותרת (סגנון כותרת, שורה ממוספרת קצרה, או שורה מודגשת קצרה)
  פותחת נושא, והשורות שאחריה משויכות למסלול לפי התווית שבתחילתן.
* **Excel**: כל גיליון עם שורת כותרת כזו; נדרש openpyxl.
* **PDF**: שכבת הטקסט, שורה אחר שורה; PDF סרוק אינו נקרא, ונאמר.

    python3 scripts/read_playbook_doc.py "<המסמך>" [--out <תיקייה זמנית>/playbook-read.json]

קוד יציאה: **0** נקרא, **3** לא נמצא בו אף נושא (המסמך אינו פלייבוק, או שאינו
קריא), **2** קלט חסר, **1** `--out` בתיקיית הלקוח.
"""
import argparse
import io
import json
import os
import re
import sys

import areo_lib  # noqa: F401  מוסיף את areolegal/lib לנתיב הייבוא

TRACK_LABELS = {
    "green": r"^(?:ירוק|מסלול ירוק|עמדה מועדפת|ברירת מחדל|העמדה שלנו|green|preferred|standard position)\b",
    "yellow": r"^(?:צהוב|מסלול צהוב|נסיגה|עמדת נסיגה|ויתור מותר|פשרה|yellow|fallback|acceptable)\b",
    "red": r"^(?:אדום|מסלול אדום|קו אדום|לא מקובל|red|red line|unacceptable|walk.?away)\b",
}
FIELD_LABELS = {
    "red_flags": r"^(?:סימני אדום|סימני זיהוי|סימנים|red flags|flags|watch for)\b",
    "alternatives": r"^(?:נוסח חלופי|נוסח מוצע|נוסח|alternative wording|proposed wording|wording|clause)\b",
    "segment": r"^(?:סגמנט|פלח|segment)\b",
    "notes": r"^(?:הערה|הערות|הסבר|נימוק|note|notes|rationale)\b",
}
OPERATING = re.compile(r"הוראות הפעלה|כללי הפעלה|אופן השימוש|operating (?:rules|instructions)|how to use", re.I)
HEADER_TRACKS = {"green": re.compile(r"ירוק|green|מועדפ|preferred", re.I),
                 "yellow": re.compile(r"צהוב|yellow|נסיגה|fallback", re.I),
                 "red": re.compile(r"אדום|red|קו אדום", re.I)}
HEADER_TOPIC = re.compile(r"נושא|סעיף|topic|clause|issue", re.I)
HEADER_FLAGS = re.compile(r"סימני|flags", re.I)


def strip_label(text, rx):
    return re.sub(rx + r"\s*[:\-.]?\s*", "", text, count=1, flags=re.I).strip()


def topic_from_table(rows):
    """rows: רשימת רשימות. מחזיר נושאים כשבשורת הכותרת יש עמודות של מסלולים."""
    if not rows:
        return []
    head = [str(c or "") for c in rows[0]]
    cols = {k: next((i for i, h in enumerate(head) if rx.search(h) and not HEADER_FLAGS.search(h)), None)
            for k, rx in HEADER_TRACKS.items()}
    if sum(v is not None for v in cols.values()) < 2:
        return []
    tcol = next((i for i, h in enumerate(head) if HEADER_TOPIC.search(h)), 0)
    fcol = next((i for i, h in enumerate(head) if HEADER_FLAGS.search(h)), None)
    out = []
    for r in rows[1:]:
        r = [str(c or "").strip() for c in r]
        if not any(r) or tcol >= len(r) or not r[tcol]:
            continue
        t = {"topic": r[tcol], "red_flags": [], "alternatives": [], "notes": []}
        for k, i in cols.items():
            t[k] = r[i] if i is not None and i < len(r) else ""
        if fcol is not None and fcol < len(r) and r[fcol]:
            t["red_flags"] = [x.strip() for x in re.split(r"[;\n•]", r[fcol]) if x.strip()]
        out.append(t)
    return out


def topics_from_lines(lines):
    """lines: [(text, is_heading)]. כותרת פותחת נושא; שורה עם תווית משויכת למסלול או לשדה."""
    topics, segments, operating, cur, in_ops = [], [], [], None, False
    for text, heading in lines:
        text = text.strip()
        if not text:
            continue
        if heading:
            in_ops = bool(OPERATING.search(text))
            if not in_ops:
                cur = {"topic": re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", text), "green": "", "yellow": "", "red": "",
                       "red_flags": [], "alternatives": [], "notes": []}
                topics.append(cur)
            continue
        if in_ops:
            operating.append(text)
            continue
        if cur is None:
            continue
        for k, rx in TRACK_LABELS.items():
            if re.search(rx, text, re.I):
                cur[k] = (cur[k] + " " + strip_label(text, rx)).strip()
                break
        else:
            for k, rx in FIELD_LABELS.items():
                if re.search(rx, text, re.I):
                    val = strip_label(text, rx)
                    if k == "segment":
                        segments.append({"topic": cur["topic"], "segment": val})
                    elif k == "red_flags":
                        cur["red_flags"] += [x.strip() for x in re.split(r"[;•]", val) if x.strip()]
                    else:
                        cur[k].append(val)
                    break
            else:
                cur["notes"].append(text)
    topics = [t for t in topics if t["green"] or t["yellow"] or t["red"] or t["red_flags"]]
    return topics, segments, operating


def read_docx(path):
    import docx
    d = docx.Document(path)
    topics = []
    for tb in d.tables:
        topics += topic_from_table([[c.text for c in row.cells] for row in tb.rows])
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        style = (p.style.name or "").lower() if p.style is not None else ""
        bold = bool(p.runs) and all(r.bold for r in p.runs if r.text.strip())
        heading = style.startswith(("heading", "כותרת", "title")) or \
            (len(t) <= 70 and (bold or re.match(r"^\d+(?:\.\d+)?[.)]\s+\S", t)) and not re.search(r"[.:]$", t)
             and not any(re.search(rx, t, re.I) for rx in list(TRACK_LABELS.values()) + list(FIELD_LABELS.values())))
        lines.append((t, heading))
    more, segments, operating = topics_from_lines(lines)
    return topics + more, segments, operating


def read_pdf(path):
    from risk.page_text import PdfTextUnavailable, pdf_text
    with open(path, "rb") as fh:
        raw = fh.read()
    try:
        text = pdf_text(raw)
    except PdfTextUnavailable as exc:
        raise ValueError("ה-PDF אינו נקרא (%s): בקשו מהלקוח עותק Word" % exc)
    lines = []
    for l in text.splitlines():
        t = l.strip()
        heading = bool(re.match(r"^\d+(?:\.\d+)?[.)]?\s+\S", t)) and len(t) <= 70 and not re.search(r"[.:]$", t)
        lines.append((t, heading))
    return topics_from_lines(lines)


def read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        raise ValueError("לקריאת Excel נדרש openpyxl (pip install openpyxl)")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    topics, operating = [], []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        found = topic_from_table(rows)
        if found:
            topics += found
        elif OPERATING.search(ws.title or ""):
            operating += [" ".join(str(c) for c in r if c) for r in rows if any(r)]
    return topics, [], operating


def main(argv=None):
    sys.dont_write_bytecode = True
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("file")
    ap.add_argument("--out", help="קובץ JSON בתיקייה זמנית; לעולם לא בתיקיית הלקוח")
    a = ap.parse_args(argv)
    if not os.path.exists(a.file):
        print("המסמך אינו קיים: %s" % a.file)
        return 2
    if a.out and os.path.commonpath([os.path.realpath(a.out), os.path.realpath(os.path.dirname(a.file))]) == \
            os.path.realpath(os.path.dirname(a.file)):
        print("--out בתיקיית הלקוח (%s). הקורא אינו כותב לשם; בחרו תיקייה זמנית." % os.path.dirname(a.file))
        return 1
    ext = os.path.splitext(a.file)[1].lower()
    try:
        if ext == ".docx":
            topics, segments, operating = read_docx(a.file)
        elif ext == ".pdf":
            topics, segments, operating = read_pdf(a.file)
        elif ext in (".xlsx", ".xlsm"):
            topics, segments, operating = read_xlsx(a.file)
        else:
            print("פורמט שאינו נקרא: %s. Word, PDF או Excel." % ext)
            return 2
    except ValueError as exc:
        print(str(exc))
        return 3
    out = {"source": os.path.basename(a.file), "format": ext[1:], "topics": topics, "segments": segments,
           "operating_rules": operating}
    if a.out:
        with io.open(a.out, "w", encoding="utf-8") as fh:
            json.dump(out, fh, ensure_ascii=False, indent=1)
    print(json.dumps(out, ensure_ascii=False, indent=1))
    if not topics:
        print("לא נמצא במסמך אף נושא עם מסלולים. ייתכן שאינו פלייבוק, או שהוא מסמך רקע (overview.md, \"מסמך רקע אינו מועמד\").")
        return 3
    return 0


if __name__ == "__main__":
    sys.exit(main())

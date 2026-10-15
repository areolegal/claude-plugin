# -*- coding: utf-8 -*-
"""בדיקות לגיליון העמדות (playbook/approval_sheet.py) ולקליטה (playbook/ingest_approval.py), לפי הערות 5.4 ו-5.5."""
import datetime as dt
import json
import os
import re
import shutil
import zipfile

import pytest
from docx import Document

import pb_helpers as H
from common import paths as PT
from common import rtl
from playbook import after_ingest as AI
from playbook import approval_sheet as A
from playbook import ingest_approval as I
from playbook import status as S

CHOICE_R22 = 'chk:choice:R-22:Q-02:0'
ASSUME_R20 = 'chk:assume:R-20:green:0'
LEGAL = {'R-22': ('עו"ד פלונית אלמונית', '14.10.2026'), 'R-15': ('עו"ד פלונית אלמונית', '14.10.2026')}


def _all_text(path):
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out += [p.text for p in c.paragraphs]
    return '\n'.join(out)


@pytest.fixture()
def sheet(tmp_path):
    m = H.good_model()
    path, manifest = A.build_sheet(m, str(tmp_path / 'to_sign'), company='EXAMPLECO', token_uid='0123456789ab')
    return m, path, manifest


def _sign(path, tmp_path, name='signed.docx', **kw):
    dst = str(tmp_path / name)
    args = dict(signers=H.ALL_SIGNERS, choices=[CHOICE_R22], assumptions=[ASSUME_R20], legal_by=LEGAL,
                remove_areo_comments=True)
    args.update(kw)
    return H.sign_sheet(path, dst, **args)


def _ingest(signed, manifest, **kw):
    kw.setdefault('received_via', 'user_upload')
    kw.setdefault('upload_ref', signed)
    return I.ingest(signed, manifest, **kw)


# ------------------------------------------------------------------ הגיליון
def test_sheet_shows_exact_wording_not_summary(sheet):
    m, path, manifest = sheet
    text = _all_text(path)
    for rule in m['rules']:
        for track in A.TRACKS:
            t = rule['traffic_light'][track]
            w = A.track_wording(track, t)
            if w:
                assert w in text, (rule['rule_id'], track)
            assert manifest['cells'][A.cell_key(rule['rule_id'], track)]['wording'] == w
    r15y = [r for r in m['rules'] if r['rule_id'] == 'R-15'][0]['traffic_light']['yellow']
    assert r15y['fallback_wording']['text'] in text


def test_sheet_shows_effect_in_each_role_for_mutual(sheet):
    m, path, _ = sheet
    text = _all_text(path)
    for rule in m['rules']:
        for track in A.TRACKS:
            t = rule['traffic_light'][track]
            assert 'כשהחברה המגלה: ' + t['as_discloser_effect'] in text
            assert 'כשהחברה המקבלת: ' + t['as_recipient_effect'] in text


def test_sheet_fails_without_recipient_effect(tmp_path):
    m = H.good_model()
    m['rules'][0]['traffic_light']['yellow'].pop('as_recipient_effect')
    with pytest.raises(A.SheetError):
        A.build_sheet(m, str(tmp_path))


def test_sheet_shows_flags_and_legal_line_apart(sheet):
    m, path, manifest = sheet
    text = _all_text(path)
    assert 'הצעת Areo' in text
    paras = text.split('\n')
    assert any(p.strip() == A.HE['legal_line'].strip() for p in paras)
    assert any(p.strip() == A.HE['legal_date'].strip() for p in paras)
    assert 'טעון בדיקה משפטיתטעון' not in text.replace('\n', '')
    assert not any(p.startswith('סימונים') and 'טעון בדיקה משפטית' in p for p in paras)
    assert manifest['cells']['R-22|-|green']['legal_review']


def test_unverified_finding_only_on_its_own_point():
    """סבב 3: "3 מקורות שוק" ו"לא אומת: 1 מקורות" באותו תא, ב-6 מתוך 10 התאים שנשענים על השוק."""
    m = H.round3_model()
    idx = A.unverified_index(m)
    for rule in m['rules']:
        for track in A.TRACKS:
            t = rule['traffic_light'][track]
            flags = A.flags_for(rule, track, t, idx)
            market = [f for f in flags if 'מקורות שוק' in f]
            assert not (market and any('לא מאומת' in f or 'לא אומת' in f for f in flags)), (rule['rule_id'], track, flags)
    r02 = [r for r in m['rules'] if r['rule_id'] == 'R-02'][0]
    assert A.flags_for(r02, 'green', r02['traffic_light']['green'], idx) == ['הצעת Areo מ-3 מקורות שוק (W3, W7, W4)']
    # הממצאים של R-02 (דרגות 1 ו-1-) אינם הנקודה של אף מסלול: הם מוצגים פעם אחת, בתא הנושא
    assert [u['level'] for u in idx[1]['R-02']] == [1, -1]


def test_unverified_label_is_one_template_with_grammar():
    assert A.unverified_label(1, ['W3'], 'נמצא מקור נספר אחד') == 'ממצא לא מאומת, מקור אחד (W3): נמצא מקור נספר אחד'
    assert A.unverified_label(2, ['W5', 'W6'], 'x') == 'ממצא לא מאומת, 2 מקורות (W5, W6): x'
    assert A.unverified_label(0, [], 'x') == 'ממצא לא מאומת, לא נמצאו מקורות: x'
    assert '()' not in A.unverified_label(0, [], 'x')


def test_sheet_text_has_no_one_sources_and_shows_topic_finding_once(tmp_path):
    m = H.round3_model()
    path, _ = A.build_sheet(m, str(tmp_path), token_uid='0123456789ad', check_wording=False)
    text = _all_text(path)
    assert '1 מקורות' not in text
    by_rule = A.unverified_index(m)[1]
    assert text.count('ממצא לא מאומת בנושא') == sum(len(v) for v in by_rule.values())
    assert 'מקור אחד (W1)' in text


def test_sheet_refuses_wording_that_fails_gap_checks(tmp_path):
    """סבב 3: היועצת חתמה על נוסח שהנושא ונותן פסל אחר כך. עכשיו הגיליון אינו מופק כשנוסח נכשל."""
    m = H.good_model()
    m['rules'][0]['traffic_light']['yellow']['fallback_wording']['text'] = 'הספק רשאי לסיים את ההסכם בשל הפרה של חובת התשלום.'
    with pytest.raises(A.SheetError) as e:
        A.build_sheet(m, str(tmp_path), token_uid='0123456789ae')
    assert m['rules'][0]['rule_id'] in str(e.value) and 'GAP_DISPUTE_CARVEOUT' in str(e.value)


def test_red_column_is_cure_wording_and_walk_away_is_tagged(sheet):
    m, path, manifest = sheet
    text = _all_text(path)
    assert 'לא נחתום (קווים אדומים)' not in text and 'לא נחתום (אדום)' not in text
    assert A.HE['walk_away'] in text
    assert text.count(A.HE['walk_away']) == sum(1 for r in m['rules'] if r.get('walk_away'))
    red = Document(path).tables[0].rows[1].cells[3].paragraphs[1].text
    assert red.startswith(A.HE['cure'].strip())
    assert 'שוברי עסקה: כלל אחד (R-22)' in text
    assert 'המלצת Areo: ' + m['validation_queue'][0]['recommendation'] in text   # ההמלצה ליד הבחירה


def test_no_approval_chain_in_sheet(sheet):
    _, path, _ = sheet
    text = _all_text(path)
    for w in ('סמכויות', 'הסלמה', 'ויתור מעבר לעמודת הצהוב', 'טעון אישור של', 'מנכ"ל'):
        assert w not in text
    assert 'חתימה לפי עמודה' in text


def test_sheet_is_hebrew_rtl_and_latin_runs_stay_ltr(sheet, tmp_path):
    _, path, _ = sheet
    assert os.path.basename(path).startswith('עמדות-לאישור')
    text = _all_text(path)
    assert 'עמדות משא ומתן לאישור' in text
    for eng in ('market_built', 'critical', 'areo_proposal', 'market_standard'):
        assert eng not in text
    assert rtl.audit_docx(path, run_date=dt.date.today()) == []
    m = H.round3_model()
    p2, _ = A.build_sheet(m, str(tmp_path / 'en'), token_uid='0123456789af', check_wording=False)
    with zipfile.ZipFile(p2) as z:
        xml = z.read('word/document.xml').decode('utf-8')
    runs = re.findall(r'<w:r>(.*?)</w:r>', xml, flags=re.S)
    latin = [r for r in runs if re.search(r'<w:t[^>]*>[^<֐-׿]*The Receiving Party', r)]
    assert latin and all('<w:rtl w:val="0"/>' in r for r in latin)
    assert not any('Receiving Party' in r and re.search('[֐-׿]', ''.join(re.findall(r'<w:t[^>]*>([^<]*)', r)))
                   for r in runs)
    assert rtl.audit_docx(p2, run_date=dt.date.today()) == []


# ------------------------------------------------------------------ הקליטה
def test_full_signature_closes_questions_and_approves(sheet, tmp_path):
    m, path, manifest = sheet
    signed = _sign(path, tmp_path)
    res = _ingest(signed, manifest, today='2026-10-14')
    assert set(v['status'] for v in res['cells'].values()) == {'approved'}
    assert res['record']['choices']['Q-02']['index'] == 0
    assert res['record']['choices']['Q-02']['text'] == m['validation_queue'][0]['options'][0]
    # לפני שהשאלה נסגרה במודל, תאי R-22 אינם נספרים: שאלה מכריעה פתוחה
    assert S.derive_status(m, res['cells'])['playbook_status'] == 'PartiallyApproved'
    rec = I.save_record(str(tmp_path / 'records'), signed, res)
    out = AI.apply_ingest(m, None, None, res, rec)
    assert out['closed'] == ['Q-02']
    q = m['validation_queue'][0]
    assert q['status'] == 'closed' and q['answered'] and q['answered_via'] == 'APR-0001'
    assert q['answer_verbatim'] == q['options'][0]
    assert out['status']['label'] == 'מדיניות מאושרת, 14.10.2026, בידי דנה לוי, יועצת משפטית ראשית'
    assert out['problems'] == []


def test_record_stores_assumptions_confirmed(sheet, tmp_path):
    m, path, manifest = sheet
    signed = _sign(path, tmp_path)
    res = _ingest(signed, manifest, today='2026-10-14')
    a = res['cells']['R-20|-|green']['assumptions_confirmed']
    assert a and a[0]['text'] == m['rules'][2]['traffic_light']['green']['assumptions'][0]
    assert a[0]['confirmed_by'] == 'דנה לוי' and a[0]['at'] == '2026-10-14'
    rec = I.save_record(str(tmp_path / 'records'), signed, res)
    AI.apply_ingest(m, None, None, res, rec)
    entry = m['approval_records'][-1]
    assert entry['assumptions_confirmed'][0]['cell'] == 'R-20|-|green'
    assert entry['choices']['Q-02']['index'] == 0
    with open(rec, encoding='utf-8') as fh:
        saved = json.load(fh)
    assert saved['choices']['Q-02']['tag'] == CHOICE_R22


def test_gate_allows_only_the_approved_wording(sheet, tmp_path):
    m, path, manifest = sheet
    res = _ingest(_sign(path, tmp_path), manifest, today='2026-10-14')
    cell = manifest['cells']['R-07|-|green']
    same = A.fingerprint([cell['position'], cell['wording']] + list(cell['effects'].values()))
    assert S.policy_claim(m, res['cells'], 'R-07', 'green', wording_fp=same)['policy_claim']
    other = A.fingerprint([cell['position'], cell['wording'] + ' בשינוי קטן'] + list(cell['effects'].values()))
    g = S.policy_claim(m, res['cells'], 'R-07', 'green', wording_fp=other)
    assert not g['policy_claim'] and g['reason'] == 'wording_differs_from_approved'


def test_open_areo_comment_blocks_the_row(sheet, tmp_path):
    m, path, manifest = sheet
    signed = _sign(path, tmp_path, remove_areo_comments=False)
    res = _ingest(signed, manifest)
    for track in A.TRACKS:
        c = res['cells']['R-22|-|%s' % track]
        assert c['status'] == 'pending_choice' and not c['policy_claim']
        assert 'R-22' in c['reason']
    assert res['cells']['R-07|-|green']['status'] == 'approved'
    assert S.derive_status(m, res['cells'])['playbook_status'] == 'PartiallyApproved'


def test_unresolved_or_double_choice_blocks_the_row(sheet, tmp_path):
    _, path, manifest = sheet
    res = _ingest(_sign(path, tmp_path, choices=[]), manifest)
    for track in A.TRACKS:
        assert res['cells']['R-22|-|%s' % track]['status'] == 'pending_choice'
        assert 'לא נבחרה אפשרות' in res['cells']['R-22|-|%s' % track]['reason']
    both = [CHOICE_R22, 'chk:choice:R-22:Q-02:1']
    res2 = _ingest(_sign(path, tmp_path, name='s2.docx', choices=both), manifest)
    assert 'יותר מאפשרות אחת' in res2['cells']['R-22|-|green']['reason']
    assert 'Q-02' not in res2['record']['choices']


def test_unconfirmed_assumption_blocks_the_cell(sheet, tmp_path):
    _, path, manifest = sheet
    res = _ingest(_sign(path, tmp_path, assumptions=[]), manifest)
    assert res['cells']['R-20|-|green']['status'] == 'pending_choice'
    assert res['cells']['R-20|-|yellow']['status'] == 'approved'


def test_legal_review_needs_name_and_date(sheet, tmp_path):
    _, path, manifest = sheet
    res = _ingest(_sign(path, tmp_path, legal_by={}), manifest)
    assert res['cells']['R-22|-|green']['status'] == 'pending_legal_review'
    assert not res['cells']['R-22|-|green']['policy_claim']
    res2 = _ingest(_sign(path, tmp_path, name='s2.docx', legal_by={'R-22': ('עו"ד פלונית אלמונית', ''),
                                                                    'R-15': LEGAL['R-15']}), manifest)
    c = res2['cells']['R-22|-|green']
    assert c['status'] == 'pending_legal_review' and c['legal_review_done_by'] == 'עו"ד פלונית אלמונית'
    res3 = _ingest(_sign(path, tmp_path, name='s3.docx'), manifest)
    assert res3['cells']['R-22|-|green']['legal_review_done_at'] == '14.10.2026'


def test_role_name_is_refused_and_no_authority_is_required(sheet, tmp_path):
    _, path, manifest = sheet
    role_only = dict(H.ALL_SIGNERS, red=('היועצת המשפטית של EXAMPLECO', 'היועצת המשפטית', '14.10.2026'))
    res = _ingest(_sign(path, tmp_path, signers=role_only), manifest)
    assert res['cells']['R-07|-|red']['status'] == 'not_signed'
    assert res['cells']['R-07|-|yellow']['status'] == 'approved'   # בלי "מאשר הסלמה": אין שרשרת אישורים
    assert res['cells']['R-07|-|green']['status'] == 'approved'
    assert 'authority' not in res['record']


def test_evidence_needs_upload_event(sheet, tmp_path):
    m, path, manifest = sheet
    signed = _sign(path, tmp_path)
    for kw in ({'received_via': None}, {'received_via': 'email'}, {'received_via': 'chat_attachment', 'upload_ref': ''},
               {'received_via': 'chat_attachment', 'upload_ref': path}):
        res = _ingest(signed, manifest, **kw)
        assert res['record']['simulated'] is True, kw
        assert not any(v['policy_claim'] for v in res['cells'].values())
    ok = _ingest(signed, manifest, received_via='chat_attachment', upload_ref='att_7f3a9c21')
    assert ok['record']['simulated'] is False
    # קובץ שזהה בייט לבייט לגיליון שהופק אינו חתום
    sim, warn = I.evidence_status('user_upload', path, manifest['issued_sha256'], manifest['issued_sha256'])
    assert sim and 'זהה בייט לבייט' in warn[0]


def test_reingest_of_the_stored_copy_does_not_fail(sheet, tmp_path):
    """סבב 3: save_record נכשל ב-SameFileError כשהקובץ החתום כבר שמור ב-records בשם הקצר."""
    _, path, manifest = sheet
    signed = _sign(path, tmp_path)
    recs = str(tmp_path / 'records')
    res = _ingest(signed, manifest)
    p1 = I.save_record(recs, signed, res)
    stored = os.path.join(recs, res['record']['signed_copy'])
    res2 = _ingest(stored, manifest, upload_ref=signed)
    assert I.save_record(recs, stored, res2) == p1          # אותו קובץ, אותה ראיה: אותה רשומה
    assert sorted(x for x in os.listdir(recs) if x.startswith('APR-')) == ['APR-0001.json']


def test_same_file_with_other_evidence_needs_supersedes(sheet, tmp_path):
    """סבב 3: אותו קובץ נקלט כ"מדומה" (email) ואחר כך כ"מאושר" (chat_attachment), בלי שום ראיה חדשה."""
    _, path, manifest = sheet
    signed = _sign(path, tmp_path)
    recs = str(tmp_path / 'records')
    I.save_record(recs, signed, _ingest(signed, manifest, received_via='email'))
    res = _ingest(signed, manifest, received_via='chat_attachment', upload_ref='att_7f3a9c21')
    with pytest.raises(I.IngestRefused):
        I.save_record(recs, signed, res)
    p2 = I.save_record(recs, signed, res, supersedes='APR-0001', reason='הקובץ הועלה מחדש לשיחה בידי היועצת')
    assert os.path.basename(p2) == 'APR-0002.json'
    path_now, rec_now = I.current_record(recs)
    assert path_now == p2 and rec_now['supersedes'] == 'APR-0001'
    with open(os.path.join(recs, 'APR-0001.json'), encoding='utf-8') as fh:
        assert json.load(fh)['superseded_by'] == 'APR-0002'


def test_simulated_signature_never_opens_policy_gate(sheet, tmp_path):
    m, path, manifest = sheet
    res = I.ingest(_sign(path, tmp_path), manifest, received_via=None)
    assert res['record']['simulated'] is True
    assert not any(v['policy_claim'] for v in res['cells'].values())
    assert S.derive_status(m, res['cells'])['playbook_status'] == 'Draft'


def test_ingest_reads_long_paths_and_refuses_to_write_beyond_the_limit(sheet, tmp_path):
    _, path, manifest = sheet
    signed = _sign(path, tmp_path)
    deep = str(tmp_path)
    while len(deep) < 300:
        deep = os.path.join(deep, 'תיקיית-לקוחה-עם-שם-ארוך-מאוד')
    os.makedirs(I.long_path(deep), exist_ok=True)
    long_signed = os.path.join(deep, 'עמדות-לאישור-הסכם סודיות הדדי (EXAMPLECO מגלה בעיקר)-v1-חתום.docx')
    assert len(long_signed) > 300
    shutil.copyfile(signed, I.long_path(long_signed))
    res = _ingest(long_signed, manifest, upload_ref='att_7f3a9c21', today='2026-10-14')
    assert set(v['status'] for v in res['cells'].values()) == {'approved'}
    with pytest.raises(PT.PathTooLong):
        I.save_record(os.path.join(deep, 'records'), long_signed, res)
    rec_path = I.save_record(str(tmp_path / 'records'), long_signed, res)
    assert os.path.exists(rec_path)
    assert len(res['record']['signed_copy']) == len('0123456789ab.docx')


def test_signed_copy_from_word_keeps_rtl(sheet, tmp_path):
    """סבב 3: העותק החתום מעזר הבדיקה נכשל ב-audit_docx (15 ריצות בלי rtl)."""
    _, path, _ = sheet
    signed = _sign(path, tmp_path)
    assert rtl.audit_docx(signed, run_date=dt.date.today()) == []

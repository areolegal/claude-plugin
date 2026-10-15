# -*- coding: utf-8 -*-
"""הכלליות של מנגנוני הפלייבוק (הערות 5.5, החלטה 22): אותם מנגנונים, אותן בדיקות, בכל סוג הסכם.

שלוש משפחות שאינן הסכם סודיות: SaaS (החברה לקוחה, תפקיד אחד), אספקה מסגרת (החברה קונה וגם מוכרת) ושכירות
(החברה שוכרת). הסכם הסודיות ההדדי נשאר כבדיקת רגרסיה. אין כאן שום בדיקה שתלויה בתוכן של סוג הסכם."""
import copy
import os

import pytest
from docx import Document

import pb_helpers as H
from playbook import after_ingest as AI
from playbook import approval_sheet as A
from playbook import ingest_approval as I
from playbook import market_counts as MC
from playbook import status as S
from playbook import surfaces_render as R
from playbook import validate_model as V

FAMILIES = ['saas_customer_excerpt.json', 'supply_buyer_seller_excerpt.json', 'lease_tenant_excerpt.json',
            'mutual_nda_model_good.json']
GAP_TEXT = 'הספק רשאי לסיים את ההסכם בשל הפרה של חובת התשלום.'


def _m(name):
    return copy.deepcopy(H.load(name))


def _all_text(path):
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out += [p.text for p in c.paragraphs]
    return '\n'.join(out)


def _sign_everything(path, manifest, tmp_path, m):
    choices = [q['options'][0]['tag'] for q in manifest['questions'].values()]
    assumptions = [t for c in manifest['cells'].values() for t in c['assumptions']]
    legal = {r['rule_id']: ('עו"ד פלונית אלמונית', '14.10.2026') for r in m['rules'] if A.legal_review_tracks(r)}
    return H.sign_sheet(path, str(tmp_path / 'signed.docx'), H.ALL_SIGNERS, choices=choices,
                        assumptions=assumptions, legal_by=legal, remove_areo_comments=True)


@pytest.mark.parametrize('name', FAMILIES)
def test_clean_family_passes_every_check(name):
    assert V.errors(V.validate(_m(name))) == []


@pytest.mark.parametrize('name', FAMILIES)
def test_sheet_shows_one_effect_line_per_company_role(name, tmp_path):
    m = _m(name)
    path, manifest = A.build_sheet(m, str(tmp_path), token_uid='0123456789ab')
    text = _all_text(path)
    roles = A.company_roles(m)
    for rule in m['rules']:
        for track in A.TRACKS:
            t = rule['traffic_light'][track]
            if len(roles) == 1:
                assert A.SINGLE_ROLE_LABEL + A.role_effect(t, roles[0][0]) in text
            else:
                for role, label in roles:
                    assert A.ROLE_LINE.format(label=label) + A.role_effect(t, role) in text
    assert len(next(iter(manifest['cells'].values()))['effects']) == len(roles)


def test_missing_role_effect_stops_the_sheet_in_any_family(tmp_path):
    m = _m('supply_buyer_seller_excerpt.json')
    m['rules'][1]['traffic_light']['yellow']['role_effects'].pop('seller')
    assert 'ROLE_EFFECT_MISSING' in V.codes(V.validate(m), 'error')
    with pytest.raises(A.SheetError):
        A.build_sheet(m, str(tmp_path))


@pytest.mark.parametrize('name', FAMILIES)
def test_full_cycle_closes_questions_and_restamps_every_surface(name, tmp_path):
    m = _m(name)
    core = {'clauses': [{'id': r['rule_id'], 'policy_status': 'Inferred'} for r in m['rules']],
            'status_label': m['status_label'], 'playbook_status': 'Draft'}
    index_entry = {'status_label': m['status_label'], 'playbook_status': 'Draft'}
    word = H.playbook_docx(str(tmp_path / 'פלייבוק - בדיקה, פלייבוק מוצע.docx'), m, S.base_label(m))
    ws = H.workspace_html(str(tmp_path / 'ws.html'), m, S.base_label(m))
    path, manifest = A.build_sheet(m, str(tmp_path / 'sheet'), token_uid='0123456789ab')
    signed = _sign_everything(path, manifest, tmp_path, m)
    res = I.ingest(signed, manifest, received_via='chat_attachment', upload_ref=signed, today='2026-10-14')
    assert set(c['status'] for c in res['cells'].values()) == {'approved'}
    rec = I.save_record(str(tmp_path / 'records'), signed, res)
    out = AI.apply_ingest(m, core, index_entry, res, rec, word_path=word, workspace_path=ws)
    assert out['problems'] == [], out['problems']
    assert out['status']['playbook_status'] == 'Approved'
    assert all(q['status'] == 'closed' and q['answered_via'] == 'APR-0001' for q in m['validation_queue'])
    assert 'מדיניות מאושרת' in os.path.basename(out['word'])
    assert R.read_workspace_status(out['workspace'])['meta']['label_now'] == out['status']['label']
    assert I.save_record(str(tmp_path / 'records'), os.path.join(str(tmp_path / 'records'),
                                                                  res['record']['signed_copy']), res) == rec


@pytest.mark.parametrize('name', FAMILIES)
def test_unverified_finding_stays_on_its_own_point(name):
    m = _m(name)
    m['market_research']['ranges'] = [{'topic': m['rules'][0]['rule_id'],
                                       'ours': {'level': 1, 'source_ids': ['W1', 'W2', 'W3']}}]
    m['market_research']['unverified_findings'] = [
        {'rule_id': m['rules'][0]['rule_id'], 'level': -1, 'source_ids': ['W4'], 'sources_found': 1, 'reason': 'x'}]
    m['rules'][0]['traffic_light']['green'].update(basis='market_standard', market_source_ids=['W1', 'W2', 'W3'],
                                                   range_point='ours')
    idx = A.unverified_index(m)
    g = m['rules'][0]['traffic_light']['green']
    flags = A.flags_for(m['rules'][0], 'green', g, idx)
    assert flags == ['הצעת Areo מ-3 מקורות שוק (W1, W2, W3)']
    assert len(idx[1][m['rules'][0]['rule_id']]) == 1


@pytest.mark.parametrize('name', FAMILIES)
def test_generic_defects_are_caught_in_every_family(name, tmp_path):
    base = _m(name)
    cases = []
    m = copy.deepcopy(base)
    low = [r for r in m['rules'] if str(r.get('importance')) != 'critical'][0]
    low['walk_away'], low['walk_away_source'] = True, 'client_answer'
    cases.append((m, 'WALKAWAY_NOT_CRITICAL'))
    m = copy.deepcopy(base)
    m['authority'] = {'escalation': 'מנכ"ל'}
    cases.append((m, 'APPROVER_FIELD'))
    m = copy.deepcopy(base)
    m['presentation_label'] = m['status_label']
    cases.append((m, 'LABEL_DUAL_FIELD'))
    m = copy.deepcopy(base)
    m['market_research']['disclosure']['fixed_sentence'] = 'הפלייבוק נבנה לפי 99 הסכמים שפורסמו ברשת.'
    cases.append((m, 'MARKET_COUNT_MISMATCH'))
    m = copy.deepcopy(base)
    m['rules'][0]['traffic_light']['yellow']['fallback_wording']['text'] = GAP_TEXT
    cases.append((m, 'WORDING_GAP'))
    for model, code in cases:
        assert code in V.codes(V.validate(model), 'error'), (name, code)
    with pytest.raises(A.SheetError):
        A.build_sheet(cases[-1][0], str(tmp_path))


@pytest.mark.parametrize('name', FAMILIES)
def test_counts_sentence_is_derived_in_every_family(name):
    m = _m(name)
    s = MC.basis_sentence(m)
    assert s and MC.count_claims(s, m) == []
    assert MC.count_claims(m['market_research']['disclosure']['fixed_sentence'], m) == []

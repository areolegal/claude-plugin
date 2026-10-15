# -*- coding: utf-8 -*-
"""בדיקות ל-common/docx_tracked.py, על הטיוטה האמיתית של C1ב מסבב 2 (fixture) ועל טיוטות קטנות.

מוכיחות: "דחה הכול" מחזיר את הטיוטה תו בתו; "קבל הכול" נותן את נוסחי היעד; XML תקין ש-Word פותח
(מרחבי השמות נשמרים, וכל קידומת ב-mc:Ignorable מוצהרת); trackRevisions דלוק במקומו; תאריך אמיתי עם dateUtc;
אין צבע; הבועות הפנימיות והכותרת הפנימית מוסרות בעותק למשלוח; ובמחיקה ובהוספה של הפסקה האחרונה אין פסקה ריקה.
"""
import datetime as _dt
import os
import re
import subprocess
import zipfile

import pytest
from lxml import etree

import orch_helpers as H
from common import docx_tracked as dt

W = dt.W


@pytest.fixture(scope='module')
def built(tmp_path_factory):
    out = str(tmp_path_factory.mktemp('c1b') / 'work.docx')
    dec = H.load_decisions('c1b_decisions_round3_clean.json')
    res = dt.apply_decisions(H.DRAFT, dec, out_path=out)
    assert res['ok'], (res['not_marked'], res['verify'])
    return out, res, dec


def test_reject_all_reproduces_received_text_exactly(built):
    out, res, _ = built
    assert dt.reject_all(out) == dt.received_paragraphs(H.DRAFT)


def test_accept_all_yields_targets(built):
    out, res, dec = built
    acc = dt.accept_all(out)
    assert acc == res['expected_accept']
    for d in dec['decisions']:
        if d['op'] == 'replace':
            assert d['new_text'] in acc
        elif d['op'] == 'replace_span':
            assert any(d['new'] in p for p in acc)
        elif d['op'] == 'insert_paragraph':
            assert d.get('number', '') + d['text'] in acc
        elif d['op'] == 'delete_paragraph':
            assert dt.received_paragraphs(H.DRAFT)[d['para']] not in acc


def test_untouched_paragraphs_identical_and_edits_are_surgical(built):
    out, res, dec = built
    root = dt.document_root(out)
    touched = {d.get('para') for d in dec['decisions']}
    received = dt.received_paragraphs(H.DRAFT)
    paras = [p for p in dt.body_paragraphs(root) if dt._mark(p) != 'ins']
    for i, p in enumerate(paras):
        if i not in touched:
            assert not list(p.iter(W + 'ins')) and not list(p.iter(W + 'del')), i
    # החלפת טווח בפסקת התשלום אינה מוחקת את כל הפסקה
    deleted = dt.revision_text(paras[18], 'del')
    assert 0 < len(deleted) < len(received[18]) / 2


def test_namespaces_preserved_and_ignorable_declared(built):
    out, _, _ = built
    orig = etree.fromstring(zipfile.ZipFile(H.DRAFT).read('word/document.xml'))
    z = zipfile.ZipFile(out)
    root = etree.fromstring(z.read('word/document.xml'))
    for prefix, uri in orig.nsmap.items():
        assert root.nsmap.get(prefix) == uri, prefix
    ign = root.get(dt.MC_IGNORABLE).split()
    assert 'w14' in ign and 'w16du' in ign
    assert all(p in root.nsmap for p in ign)
    for name in z.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):
            r = etree.fromstring(z.read(name))  # כל חלק הוא XML תקין
            assert all(p in r.nsmap for p in (r.get(dt.MC_IGNORABLE) or '').split()), name
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
    for target in re.findall(r'Target="([^"]+)"', rels):
        if not target.startswith('http') and 'TargetMode' not in target:
            assert 'word/' + target in z.namelist() or target.startswith('../'), target
    ct = z.read('[Content_Types].xml').decode('utf-8')
    assert '/word/comments.xml' in ct and '/word/headerInternal.xml' in ct


def test_track_revisions_on_in_schema_order(built):
    out, _, _ = built
    st = etree.fromstring(zipfile.ZipFile(out).read('word/settings.xml'))
    assert st.find(W + 'trackRevisions') is not None
    assert st.find(W + 'trackChanges') is None
    kids = [dt.local(k.tag) for k in st]
    before = kids[:kids.index('trackRevisions')]
    assert all(k in dt.SETTINGS_BEFORE_TRACK for k in before)


def test_real_dates_with_dateutc(built):
    out, _, _ = built
    root = dt.document_root(out)
    revs = [e for e in root.iter(W + 'ins', W + 'del') if e.get(W + 'author') == H.AUTHOR]
    assert revs
    now = _dt.datetime.now()
    offset = _dt.datetime.now().astimezone().utcoffset()
    for e in revs:
        d, u = e.get(W + 'date'), e.get(dt.DATE_UTC_ATTR)
        assert re.fullmatch(r'\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ', d)
        assert re.fullmatch(r'\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ', u)
        loc = _dt.datetime.strptime(d, dt.DATE_FMT)
        assert now - _dt.timedelta(hours=1) <= loc <= now + _dt.timedelta(minutes=1)
        assert loc - _dt.datetime.strptime(u, dt.DATE_FMT) == offset
    assert len({e.get(W + 'date') for e in revs}) > 1  # לא כל השינויים באותה שנייה


def test_no_highlight_added(built, tmp_path):
    out, _, _ = built
    assert zipfile.ZipFile(out).read('word/document.xml').count(b'w:highlight') == \
        zipfile.ZipFile(H.DRAFT).read('word/document.xml').count(b'w:highlight')
    src = H.small_draft(str(tmp_path / 'hl.docx'), ['טקסט מסומן בצבע אצלם.'], highlight_first_run=True)
    doc = dt.TrackedDocx(src, H.AUTHOR, H.INITIALS)
    doc.replace_span(0, 'אצלם', 'אצל הצד שכנגד', external='תיקון מינוח.')
    doc.save(str(tmp_path / 'hl_out.docx'))
    root = dt.document_root(str(tmp_path / 'hl_out.docx'))
    for ins in root.iter(W + 'ins'):
        assert not list(ins.iter(W + 'highlight'))
    assert doc.verify_roundtrip() == []


def test_rtl_runs_split_by_script(built):
    out, _, _ = built
    root = dt.document_root(out)
    for ins in root.iter(W + 'ins'):
        if dt.local(ins.getparent().tag) == 'rPr':
            continue
        for r in ins.iter(W + 'r'):
            txt = ''.join(t.text or '' for t in r.iter(W + 't'))
            rtl = r.find(W + 'rPr/' + W + 'rtl') is not None
            if re.search(r'[\u0590-\u05FF]', txt):
                assert rtl, txt
            if re.fullmatch(r'[A-Za-z][A-Za-z0-9 ./&-]*', txt.strip() or '-'):
                assert not rtl, txt


def test_internal_comments_and_header_stripped_in_send_copy(built, tmp_path):
    out, _, dec = built
    send = str(tmp_path / 'send.docx')
    dt.strip_internal(out, send)
    work_c, send_c = dt.read_comments(out), dt.read_comments(send)
    assert any(c['author'] == dt.INTERNAL_AUTHOR for c in work_c)
    assert not any(c['author'] == dt.INTERNAL_AUTHOR for c in send_c)
    assert len([c for c in send_c if c['author'] == H.AUTHOR]) == len([c for c in work_c if c['author'] == H.AUTHOR])
    z = zipfile.ZipFile(send)
    assert dt.INTERNAL_HEADER_PART not in z.namelist()
    root = etree.fromstring(z.read('word/document.xml'))
    ids_doc = {e.get(W + 'id') for e in root.iter(W + 'commentReference')}
    ids_cm = {c['id'] for c in send_c}
    assert ids_doc == ids_cm
    assert {e.get(W + 'id') for e in root.iter(W + 'commentRangeStart')} == ids_cm
    assert all(p in root.nsmap for p in root.get(dt.MC_IGNORABLE).split())
    assert dt.reject_all(send) == dt.received_paragraphs(H.DRAFT)
    assert dt.accept_all(send) == dt.accept_all(out)


def test_existing_header_restored_in_send_copy(tmp_path):
    from docx import Document
    d = Document()
    d.sections[0].header.paragraphs[0].text = 'כותרת של הצד שכנגד'
    d.add_paragraph('סעיף אחד.')
    src = str(tmp_path / 'hdr.docx')
    d.save(src)
    doc = dt.TrackedDocx(src, H.AUTHOR, H.INITIALS)
    doc.replace_span(0, 'אחד', 'ראשון', external='תיקון.')
    work = doc.save(str(tmp_path / 'hdr_work.docx'))
    send = str(tmp_path / 'hdr_send.docx')
    dt.strip_internal(work, send)
    z = zipfile.ZipFile(send)
    hdr = ''.join(z.read(n).decode('utf-8') for n in z.namelist() if n.startswith('word/header'))
    assert 'כותרת של הצד שכנגד' in hdr and dt.INTERNAL_HEADER_TEXT not in hdr
    ref = etree.fromstring(z.read('word/document.xml')).find('.//' + W + 'headerReference')
    assert not ref.get('{%s}id' % dt.R_NS).startswith(dt.INTERNAL_REL_PREFIX)


def test_metadata_cleaned(built):
    out, _, _ = built
    z = zipfile.ZipFile(out)
    core = z.read('docProps/core.xml').decode('utf-8')
    assert H.AUTHOR in core and 'python-docx' not in core
    assert 'docProps/thumbnail.jpeg' not in z.namelist()
    assert not any(n.startswith('customXml/') for n in z.namelist())


@pytest.mark.parametrize('case', ['delete_last', 'delete_last_and_add', 'add_after_last'])
def test_last_paragraph_cases_leave_no_empty_paragraph(tmp_path, case):
    src = H.small_draft(str(tmp_path / 'syn.docx'), ['הסכם סודיות', '1. הצדדים מסכימים.', '2. הסודיות תישמר.',
                                                     'לחיצה על "הסכם" מהווה אישור.'])
    doc = dt.TrackedDocx(src, H.AUTHOR, H.INITIALS)
    if case in ('delete_last', 'delete_last_and_add'):
        doc.delete_paragraph(3, external='מחקנו את פסקת הלחיצה.')
    if case in ('delete_last_and_add', 'add_after_last'):
        doc.insert_paragraph(3, 'ולראיה באו הצדדים על החתום:', external='הוספנו בלוק חתימה.')
        doc.insert_paragraph(3, 'EXAMPLECO: ' + '________ (יושלם בחתימה)')
    out = doc.save(str(tmp_path / ('%s.docx' % case)))
    assert doc.verify_roundtrip() == []
    acc, rej = dt.accept_all(out), dt.reject_all(out)
    assert rej == dt.received_paragraphs(src)
    assert acc and acc[-1].strip(), acc  # אין פסקה ריקה בסוף ההסכם
    assert all(t.strip() for t in acc)


def test_multiple_spans_in_one_paragraph_are_one_marking(tmp_path):
    doc = dt.TrackedDocx(H.DRAFT, H.AUTHOR, H.INITIALS)
    target = doc.edit_spans(3, [('[שם הלקוח]', 'EXAMPLECO בע"מ', 'השלמנו את פרטי הלקוח.', None),
                                ('("הספק" או "נותן השירותים")', '("הספק", "היועץ" או "נותן השירותים")',
                                 'הוספנו את הכינוי "היועץ".', 'פנימי: A2')])
    with pytest.raises(dt.UnsupportedEdit):
        doc.replace_span(3, 'בתל אביב', 'בתל אביב-יפו', external='x')  # פסקה מסומנת פעם אחת בלבד
    out = doc.save(str(tmp_path / 'spans.docx'))
    p = dt.body_paragraphs(dt.document_root(out))[3]
    starts = [s.get(W + 'id') for s in p.iter(W + 'commentRangeStart')]
    assert len(starts) == 3  # שתי בועות חיצוניות, כל אחת לטווח שלה, ובועה פנימית אחת
    assert dt.accept_all(out)[3] == target
    assert dt.reject_all(out) == dt.received_paragraphs(H.DRAFT)


def test_unsupported_edit_is_reported_and_not_applied(tmp_path):
    dec = H.decisions({'id': 'X1', 'op': 'replace_span', 'para': 18, 'old': 'טקסט שאינו קיים', 'new': 'משהו',
                       'decision_code': 'R1', 'external_comment': 'בועה.'})
    res = dt.apply_decisions(H.DRAFT, dec, out_path=str(tmp_path / 'x.docx'))
    assert res['not_marked'] and res['not_marked'][0]['id'] == 'X1'
    assert dt.reject_all(str(tmp_path / 'x.docx')) == dt.received_paragraphs(H.DRAFT)


def test_validate_decisions_contract():
    bad = H.decisions(
        {'id': 'A', 'op': 'replace', 'para': 1, 'new_text': 'x', 'decision_code': 'R1'},             # בלי בועה
        {'id': 'B', 'op': 'insert_paragraph', 'after_para': 1, 'text': 'x', 'decision_code': 'R1',
         'external_comment': 'y'},                                                                     # נושא חסר אינו R1
        {'id': 'C', 'op': 'replace_span', 'para': 1, 'old': 'a', 'new': 'b', 'decision_code': 'N1',
         'external_comment': 'y'},                                                                     # N1 רק להוספה
        {'id': 'D', 'op': 'move', 'para': 1, 'decision_code': 'A1'})
    errs = dt.validate_decisions(bad)
    assert any(e.startswith('A:') for e in errs)
    assert any(e.startswith('B:') for e in errs)
    assert any(e.startswith('C:') for e in errs)
    assert any(e.startswith('D:') for e in errs)
    assert dt.validate_decisions(H.load_decisions('c1b_decisions_round3_clean.json')) == []


def test_char_ops_keeps_numbers_whole_and_refines_letters():
    ops = dt.char_ops('בתוך 30 יום', 'בתוך 45 יום')
    assert ('delete', 5, 7, '') in ops and any(o[0] == 'insert' and o[3] == '45' for o in ops)
    ops = dt.char_ops('הלקוח רשאי', 'הלקוחה רשאי')
    assert [o for o in ops if o[0] != 'equal'] == [('insert', 5, 5, 'ה')]


@pytest.mark.skipif(os.environ.get('AREO_WORD_CHECK') != '1', reason='בדיקה מול Word עצמו: רק כש-AREO_WORD_CHECK=1')
def test_word_accept_reject_match_projection(built, tmp_path):  # pragma: no cover
    out, res, _ = built
    script = os.path.join(H.ROOT, 'tests', 'word_views.ps1')
    subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', script, '-InPath', out,
                    '-OutDir', str(tmp_path)], check=True, timeout=300)
    for mode, mine in (('accept', dt.accept_all(out)), ('reject', dt.reject_all(out))):
        lines = open(str(tmp_path / ('word_%s.txt' % mode)), encoding='utf-8-sig').read().split('\n')
        lines = [x.rstrip('\r') for x in lines]
        while lines and lines[-1] == '':
            lines.pop()
        assert lines == [m.replace('\n', '\x0b') for m in mine]  # Word returns \x0b for w:br; the projection has \n

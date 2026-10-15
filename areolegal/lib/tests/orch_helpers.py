# -*- coding: utf-8 -*-
"""עזרי בדיקה לרכיב הנושא ונותן: נתיבי fixtures, פרופיל לדוגמה, ובניית טיוטה קטנה ב-python-docx."""
import io
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from docx import Document  # noqa: E402

FIX = os.path.join(ROOT, 'tests', 'fixtures', 'orchestrator')
DRAFT = os.path.join(FIX, 'c1b_counterparty_draft.docx')
AUTHOR = 'עורכת הדין לדוגמה'
INITIALS = 'עד'
PROFILE = {'lawyer_name': AUTHOR, 'lawyer_initials': INITIALS, 'organization_name': 'EXAMPLECO',
           'organization_id': '511111114'}


def load_decisions(name):
    with io.open(os.path.join(FIX, name), encoding='utf-8') as fh:
        return json.load(fh)


def small_draft(path, paragraphs, highlight_first_run=False):
    """טיוטה קטנה של "הצד שכנגד", בלי סימוני מעקב."""
    doc = Document()
    for k, t in enumerate(paragraphs):
        p = doc.add_paragraph()
        r = p.add_run(t)
        if highlight_first_run and k == 0:
            from docx.enum.text import WD_COLOR_INDEX
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.save(path)
    return path


def decisions(*items, tid='t'):
    return {'transaction_id': tid, 'round': 1, 'received_sha256': 'x', 'author': AUTHOR, 'initials': INITIALS,
            'decisions': list(items)}

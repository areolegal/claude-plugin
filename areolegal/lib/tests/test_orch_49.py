# -*- coding: utf-8 -*-
"""בדיקות לשינויים של הערות הנושא ונותן 4.9 (סבב הבדיקה 3).

(1) החלטת מוצר 2 (למה בכל בועה): התגובה של C1 בסבב 3, כפי שנשלחה, נכשלת בכל אחד מהרכיבים שהסוקרים מצאו, והגרסה
    המתוקנת שלה עוברת את השער בשני המצבים.
(2) כלליות (החלטה 22): טיוטה שנייה בעברית (הסכם תחזוקה), הסכם סודיות באנגלית, הסכם אספקה, שכירות ו-SaaS. בכל
    משפחה הנוסח התקין עובר בשני המצבים, והפגום נכשל בבדיקות הכלליות (כיסוי, נימוק, טענה שקרית ומחלקות 2.22).
(3) נייר באנגלית: ריצות LTR, בועות בלי bidi, ושדה ריק לחתימה באנגלית.
(4) הלולאה שלפני המשלוח: שורות "לתיקון" ממעריך הסיכונים חוסמות את העותק למשלוח עד שנסגרו.
(5) החלטת מוצר 1 (אין שרשרת אישורים), קובץ המצב, מנת השאלות, "לאישורך", נקודות לשיחה ורשימת הבדיקה לפני חתימה.
(6) גבול נתיב אחד (230) ו-audit_docx שמקבל ריצה לטינית בלי rtl.

הטיוטות כאן אינן הקלט של סבב בדיקה: services2 ו-nda_en נכתבו לבדיקות בלבד. c1b היא הטיוטה של C1, ומשמשת כאן רק כבדיקת
רגרסיה של התגובה שנשלחה בסבב 3.
"""
import copy
import io
import json
import os
import re

import pytest
from lxml import etree

import orch_helpers as H
from common import docx_tracked as dt
from common import ids, paths, rtl
from orchestrator import (pre_send, proportionality as pr, qa_gate, question_batch as qb, review_table as rt,
                          signing_checklist as sc, state_builder as sb, state_migrate as sm, talking_points as tp)

FIX = H.FIX
C1 = H.DRAFT
S2 = os.path.join(FIX, 'services2_counterparty_draft.docx')
EN = os.path.join(FIX, 'nda_en_counterparty_draft.docx')
SIGNING = {'entity_details': {'company_number': {'status': 'signing_blank', 'answer': 'נשלים לפני החתימה'}}}


def _codes(g):
    return {f['code'] for f in g['failures']}


def _msgs(g, code):
    return ' '.join('%s@%s' % (f['message'], f['where']) for f in g['failures'] if f['code'] == code)


def _apply(tmp_path, draft, name, strict=True):
    dec = H.load_decisions(name)
    out = str(tmp_path / (name + '.docx'))
    if strict:
        res = dt.apply_decisions(draft, dec, out_path=out)
        assert not res['errors'] and not res['not_marked'], (res['errors'], res.get('not_marked'))
        return out, res['expected_accept'], dec
    doc = dt.TrackedDocx(draft, dec['author'], dec['initials'])
    assert dt.mark_decisions(doc, dec) == []
    doc.save(out)
    return out, doc.expected_accept(), dec


def _both(tmp_path, draft, name, state=None):
    out, exp, dec = _apply(tmp_path, draft, name)
    w = qa_gate.run_gate(out, draft, 'work', H.PROFILE, state or {}, exp)
    send = str(tmp_path / ('send_' + name + '.docx'))
    s = qa_gate.build_send_copy(out, draft, send, H.PROFILE, copy.deepcopy(state or SIGNING), exp)
    return w, s


# --------------------------------------------------------------------------------- (1) C1 בסבב 3
@pytest.fixture(scope='module')
def c1_as_sent(tmp_path_factory):
    tmp = tmp_path_factory.mktemp('c1r3')
    out, exp, dec = _apply(tmp, C1, 'c1r3_decisions_as_sent.json', strict=False)
    work = qa_gate.run_gate(out, C1, 'work', H.PROFILE, {}, exp)
    send = qa_gate.build_send_copy(out, C1, str(tmp / 'send.docx'), H.PROFILE, copy.deepcopy(SIGNING), exp)
    return work, send, dec


def test_round3_c1_as_sent_violates_the_decisions_contract(c1_as_sent):
    errs = dt.validate_decisions(c1_as_sent[2])
    assert any(e.startswith('H01:') for e in errs) and any(e.startswith('H02:') for e in errs)
    assert any(e.startswith('D31:') for e in errs)


@pytest.mark.parametrize('code,needle', [
    # בועת התשלום: אין "למה" למועד, לתקרה ולריבית
    ('COV_UNCOVERED_CHANGE', 'מועד התשלום@19'),
    ('BUB_WHY_MISSING', 'הגבלת אחריות@19'),
    ('BUB_WHY_MISSING', 'ריבית והצמדה@19'),
    # שיפוי בדרישה ראשונה
    ('BUB_WHY_MISSING', 'שיפוי וערבות@12'),
    # הרישיון הרחב
    ('COV_UNCOVERED_CHANGE', 'היקף הרישיון@15'),
    # תקופת אי-השידול
    ('BUB_WHY_MISSING', 'איסור שידול@11'),
    # "ידועה מראש"
    ('BUB_CLAIM_FALSE', 'ידוע מראש'),
    # כותרות הפרק ובלוק החתימה בלי בועה חיצונית
    ('COV_CHANGE_WITHOUT_BUBBLE', '@30'),
    ('COV_CHANGE_WITHOUT_BUBBLE', '@34'),
    ('COV_CHANGE_WITHOUT_BUBBLE', '@41'),
])
def test_round3_c1_as_sent_fails_bubble_why(c1_as_sent, code, needle):
    work, send, _ = c1_as_sent
    for g in (work, send):
        assert needle in _msgs(g, code), _msgs(g, code)


def test_round3_c1_as_sent_blocks_the_send_copy(c1_as_sent):
    work, send, _ = c1_as_sent
    assert not work['passed'] and not send['passed'] and send['send_copy'] is None


def test_round3_c1_corrected_passes_both_modes(tmp_path):
    w, s = _both(tmp_path, C1, 'c1r3_decisions_corrected.json')
    assert w['passed'], w['failures']
    assert s['passed'], s['failures']
    dec = H.load_decisions('c1r3_decisions_corrected.json')['decisions']
    c = pr.count_asks(dec)
    assert c['additions'] == len([d for d in dec if d['op'] == 'insert_paragraph']) - 2  # שתי הכותרות אינן בקשות


# --------------------------------------------------------------------------------- (2) טיוטות נוספות
FAMILIES = {
    # משפחה: (טיוטה, קודים שהנוסח הפגום חייב להפיל)
    'services2': ('services2_counterparty_draft.docx',
                  {'COV_UNCOVERED_CHANGE', 'COV_CHANGE_WITHOUT_BUBBLE', 'BUB_WHY_MISSING', 'BUB_CLAIM_FALSE',
                   'GAP_CONFIDENTIALITY_EXCEPTIONS', 'GAP_TERMINATION_HANDOVER'}),
    'supply': ('supply_counterparty_draft.docx',
               {'COV_CHANGE_WITHOUT_BUBBLE', 'BUB_WHY_MISSING', 'BUB_CLAIM_FALSE', 'GAP_CONFIDENTIALITY_EXCEPTIONS',
                'GAP_PAYMENT_WORDING', 'GAP_TERMINATION_HANDOVER'}),
    'lease': ('lease_counterparty_draft.docx',
              {'COV_UNCOVERED_CHANGE', 'COV_CHANGE_WITHOUT_BUBBLE', 'BUB_WHY_MISSING', 'BUB_CLAIM_FALSE',
               'GAP_DISPUTE_CARVEOUT', 'GAP_TERMINATION_HANDOVER'}),
    'saas': ('saas_counterparty_draft.docx',  # באנגלית
             {'COV_CHANGE_WITHOUT_BUBBLE', 'BUB_WHY_MISSING', 'BUB_CLAIM_FALSE', 'GAP_CONFIDENTIALITY_EXCEPTIONS',
              'GAP_DISPUTE_CARVEOUT', 'GAP_PAYMENT_WORDING', 'GAP_TERMINATION_HANDOVER'}),
}


@pytest.mark.parametrize('family', sorted(FAMILIES))
def test_family_clean_passes_both_modes(tmp_path, family):
    draft = os.path.join(FIX, FAMILIES[family][0])
    w, s = _both(tmp_path, draft, '%s_decisions_clean.json' % family)
    assert w['passed'], w['failures']
    assert s['passed'], s['failures']
    assert w['info']['lang'] == ('en' if family == 'saas' else 'he')  # SaaS נכתב באנגלית


@pytest.mark.parametrize('family', sorted(FAMILIES))
def test_family_defects_fail_generic_checks(tmp_path, family):
    draft, expected = FAMILIES[family]
    draft = os.path.join(FIX, draft)
    out, exp, _ = _apply(tmp_path, draft, '%s_decisions_defects.json' % family, strict=False)
    g = qa_gate.run_gate(out, draft, 'work', H.PROFILE, {}, exp)
    assert expected <= _codes(g), _codes(g)
    s = qa_gate.build_send_copy(out, draft, str(tmp_path / 'send.docx'), H.PROFILE, copy.deepcopy(SIGNING), exp)
    assert s['send_copy'] is None


def test_true_known_in_advance_claim_is_not_flagged(tmp_path):
    """"ידועה מראש" על תקרה שנוקבת סכום (services2, S04) אינה טענה שקרית."""
    out, exp, _ = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    g = qa_gate.run_gate(out, S2, 'work', H.PROFILE, {}, exp)
    assert 'BUB_CLAIM_FALSE' not in _codes(g)


EN_STATE = {'entity_details': {'company_number': {'status': 'signing_blank', 'answer': 'to be completed at signing'}}}


def test_english_nda_clean_passes_both_modes(tmp_path):
    w, s = _both(tmp_path, EN, 'nda_en_decisions_clean.json', state=EN_STATE)
    assert w['passed'], w['failures']
    assert s['passed'], s['failures']
    assert w['info']['lang'] == 'en'
    assert [x['text'] for x in s['info']['placeholders']['signing_blanks']] == [ids.signing_blank('en')] * 2


def test_user_decision_can_waive_a_gap_class_explicitly(tmp_path):
    """ויתור מכוון בהכרעה מפורשת של המשתמש, עם קוד הפער, אינו כשל. בתשובה כללית הוא נשאר כשל."""
    out, exp, _ = _apply(tmp_path, S2, 'services2_decisions_defects.json', strict=False)
    code = 'GAP_CONFIDENTIALITY_EXCEPTIONS'
    assert code in _codes(qa_gate.run_gate(out, S2, 'work', H.PROFILE, {}, exp))
    st = {'deliberate_concessions': [{'topic': 'סודיות', 'clause_ref': 'S-008', 'reason': 'תגובה ממוקדת',
                                      'gap_codes': [code], 'user_decision': {'answer_kind': 'explicit'}}]}
    assert code not in _codes(qa_gate.run_gate(out, S2, 'work', H.PROFILE, st, exp))
    st['deliberate_concessions'][0]['user_decision'] = {'answer_kind': 'delegated'}
    assert code in _codes(qa_gate.run_gate(out, S2, 'work', H.PROFILE, st, exp))


# --------------------------------------------------------------------------------- (3) נייר באנגלית
def test_english_paper_runs_and_bubbles_are_ltr(tmp_path):
    out, _, _ = _apply(tmp_path, EN, 'nda_en_decisions_clean.json')
    parts = dt.read_package(out)
    root = etree.fromstring(parts['word/document.xml'])
    W = dt.W
    ins_runs = [r for e in root.iter(W + 'ins') if e.get(W + 'author') == H.AUTHOR for r in e.iter(W + 'r')
                if ''.join(t.text or '' for t in r.iter(W + 't'))]
    assert ins_runs
    for r in ins_runs:
        txt = ''.join(t.text or '' for t in r.iter(W + 't'))
        has_rtl = r.find(W + 'rPr/' + W + 'rtl') is not None
        assert has_rtl == bool(re.search('[\u05d0-\u05ea]', txt)), txt  # פיסוק וספרות בלי rtl
    cx = etree.fromstring(parts[dt.COMMENTS_PART])
    ours = [c for c in cx.findall(W + 'comment') if c.get(W + 'author') == H.AUTHOR]
    assert ours and all(p.find(W + 'pPr/' + W + 'bidi') is None for c in ours for p in c.findall(W + 'p'))


def test_hebrew_paper_still_rtl(tmp_path):
    out, _, _ = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    root = dt.document_root(out)
    W = dt.W
    for e in root.iter(W + 'ins'):
        for r in e.iter(W + 'r'):
            txt = ''.join(t.text or '' for t in r.iter(W + 't'))
            if re.search('[\u05d0-\u05ea]', txt):
                assert r.find(W + 'rPr/' + W + 'rtl') is not None


def test_english_signing_blank_and_author_name():
    assert ids.signing_blank('en') == '________ (to be completed at signing)'
    assert ids.detail_for_draft('company_number', answer='to be completed at signing', lang='en') == \
        (ids.signing_blank('en'), 'signing_blank')
    assert qa_gate.author_for({'lawyer_name': 'עורכת', 'lawyer_name_en': 'Counsel'}, 'en') == 'Counsel'
    assert qa_gate.author_for({'lawyer_name': 'עורכת'}, 'en') == 'עורכת'
    # ראשי תיבות: מהפרופיל, ואחרת מהשם (סבב 3: בפרופיל לא היה lawyer_initials)
    assert qa_gate.initials_for({'lawyer_name': 'x', 'lawyer_name_en': 'Dana Levi'}, 'en') == 'DL'
    assert qa_gate.initials_for({'lawyer_name': 'x', 'lawyer_initials': 'QA'}) == 'QA'


def test_rtl_audit_accepts_latin_run_without_rtl(tmp_path):
    """rtl.audit_docx החזיר התרעת שווא על "EXAMPLECO" שהמנוע כתב בלי w:rtl (C1 בסבב 3, פער 2)."""
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    rtl.set_run_rtl(p.add_run('הלקוח: '))
    p.add_run('EXAMPLECO 2026')  # ריצה לטינית בלי rtl, כמו שהמנוע כותב
    rtl._ppr_rtl(p._p)  # bidi לפסקה, בלי לגעת בריצה הלטינית
    for s in d.sections:
        rtl.set_section_rtl(s)
    rtl.clean_metadata(d)
    path = str(tmp_path / 'latin.docx')
    d.save(path)
    probs = rtl.audit_docx(path)
    assert not any('EXAMPLECO' in x for x in probs), probs
    d2 = Document()
    d2.add_paragraph().add_run('טקסט עברי בלי rtl')
    p2 = str(tmp_path / 'heb.docx')
    d2.save(p2)
    assert any('ריצה בלי rtl' in x for x in rtl.audit_docx(p2))


def test_clean_metadata_writes_real_utc(tmp_path):
    """ממצא 14 של המהנדס: clean_metadata כתב שעה מקומית עם "Z". עכשיו core.xml נושא UTC אמיתי."""
    import datetime as _dt
    import zipfile
    from docx import Document
    tz = _dt.timezone(_dt.timedelta(hours=3))
    d = Document()
    rtl.clean_metadata(d, when=_dt.datetime(2026, 10, 14, 17, 20, tzinfo=tz))
    path = str(tmp_path / 'meta.docx')
    d.save(path)
    core = zipfile.ZipFile(path).read('docProps/core.xml').decode('utf-8')
    assert '2026-10-14T14:20:00Z' in core and '17:20:00Z' not in core
    before = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    d2 = Document()
    rtl.clean_metadata(d2)
    after = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None)
    assert before <= d2.core_properties.modified.replace(tzinfo=None) <= after
    naive_local = _dt.datetime(2026, 10, 14, 17, 20)
    assert rtl.utc_naive(naive_local) == naive_local.astimezone().astimezone(_dt.timezone.utc).replace(tzinfo=None)


# --------------------------------------------------------------------------------- (4) לולאה לפני משלוח
def _register(status='open', clause='S04', kind='created_by_response'):
    return {'pre_send_fixes': [{'id': 'PSF-1', 'kind': kind, 'gap': 'בסיס התקרה אינו מוגדר', 'where': 'אחריות',
                                'risk_ids': ['RSK-001'], 'clause_ref': clause, 'status': status}]}


def test_open_pre_send_fix_blocks_send_and_retracts_old_copy(tmp_path):
    out, exp, _ = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    deal = str(tmp_path / 'deal')
    os.makedirs(deal)
    send = os.path.join(deal, 'הסכם תחזוקה - EXAMPLECO - סבב 1 - 14.10.2026.docx')
    first = qa_gate.build_send_copy(out, S2, send, H.PROFILE, copy.deepcopy(SIGNING), exp, deal_dir=deal)
    assert first['passed'] and os.path.exists(send)  # אין מרשם: העותק מופק
    pre_send.save_register(deal, _register())         # מעריך הסיכונים רץ אחרי ההפקה
    st = copy.deepcopy(SIGNING)
    res = qa_gate.build_send_copy(out, S2, send, H.PROFILE, st, exp, deal_dir=deal)
    assert 'RISK_FIX_OPEN' in _codes(res) and res['send_copy'] is None
    assert not os.path.exists(send) and res['retracted'] and os.path.exists(res['retracted'])
    assert 'לא למשלוח' in os.path.basename(res['retracted'])
    w = qa_gate.run_gate(out, S2, 'work', H.PROFILE, st, exp, register=pre_send.load_register(deal))
    assert w['passed'] and any(x['code'] == 'RISK_FIX_OPEN' for x in w['warnings'])


def test_pre_send_fix_closed_by_reopened_decision(tmp_path):
    out, exp, dec = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    dec = copy.deepcopy(dec)
    next(d for d in dec['decisions'] if d['id'] == 'S04')['reopened_from_fix'] = 'PSF-1'
    reg = _register()
    pre_send.resolve_fix(reg, 'PSF-1', 'S04')
    st = copy.deepcopy(SIGNING)
    g = qa_gate.run_gate(out, S2, 'send', H.PROFILE, st, exp, register=reg, decisions=dec)
    assert 'RISK_FIX_RESOLUTION_UNVERIFIED' in _codes(g)  # סומן resolved בלי הכרעה בקובץ המצב
    st['clauses'] = [{'clause_id': 'S04', 'reopened_from_fix': 'PSF-1'}]
    assert not {'RISK_FIX_OPEN', 'RISK_FIX_RESOLUTION_UNVERIFIED'} & _codes(
        qa_gate.run_gate(out, S2, 'send', H.PROFILE, st, exp, register=reg, decisions=dec))
    # 1.2.0: הסימון לבדו אינו מספיק. בלי קובץ ההכרעות, או עם מזהה הכרעה שאינו בו, השורה אינה מאומתת
    assert 'RISK_FIX_RESOLUTION_UNVERIFIED' in _codes(
        qa_gate.run_gate(out, S2, 'send', H.PROFILE, st, exp, register=reg))
    reg2 = _register()
    pre_send.resolve_fix(reg2, 'PSF-1', 'S04b')
    st['clauses'].append({'clause_id': 'S04b', 'reopened_from_fix': 'PSF-1'})
    assert 'RISK_FIX_RESOLUTION_UNVERIFIED' in _codes(
        qa_gate.run_gate(out, S2, 'send', H.PROFILE, st, exp, register=reg2, decisions=dec))
    with pytest.raises(ValueError):
        pre_send.resolve_fix(_register(), 'PSF-1', 'S04b', decisions=dec)


def test_pre_send_fix_closed_by_deliberate_concession(tmp_path):
    out, exp, _ = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    reg = _register()
    st = copy.deepcopy(SIGNING)
    with pytest.raises(ValueError):
        pre_send.concede_fix(reg, st, 'PSF-1', 'הסיכון מקובל', {})
    pre_send.concede_fix(reg, st, 'PSF-1', 'הסיכון מקובל', {'answer_kind': 'explicit', 'answer': 'מקבלים'})
    assert reg['pre_send_fixes'][0]['status'] == 'conceded'
    assert st['deliberate_concessions'][0]['clause_ref'] == 'S04' and \
        st['deliberate_concessions'][0]['from_pre_send_fix'] == 'PSF-1'
    g = qa_gate.run_gate(out, S2, 'send', H.PROFILE, st, exp, register=reg)
    assert not {c for c in _codes(g) if c.startswith('RISK_FIX')}
    handed = pr.waiver_handoff(st)
    assert handed[0]['from_pre_send_fix'] == 'PSF-1' and handed[0]['report_as'] == 'accepted_risk'


def test_polish_fix_warns_only_and_old_field_name_is_read(tmp_path):
    reg = {'pre_send_fixes': [{'kind': 'polish', 'gap': 'ליטוש', 'state_clause_id': 'S03'}]}
    assert pre_send.check(reg) == [] and pre_send.polish_warnings(reg)
    reg2 = {'pre_send_fixes': [{'kind': 'incomplete', 'gap': 'x', 'state_clause_id': 'D22'}]}
    assert pre_send.to_reopen(reg2) == [{'fix_id': 'PSF-01', 'clause_ref': 'D22', 'gap': 'x', 'risk_ids': []}]


# --------------------------------------------------------------------------------- (5) מצב, אישורים ומודולים
def test_validate_state_rejects_round3_structures():
    st = sb.skeleton()
    st['meta'].update(round_stage='response_produced_send_blocked', round_number=1)
    st['clauses'] = [{'clause_id': 'E-013', 'policy_gate': {'claim': True},
                      'approval': {'status': 'proposed', 'required_by': 'מנכ"ל'}}]
    st['deliberate_concessions'] = [{'topic': 'דין ושיפוט', 'clause_ref': 'R-14', 'reason': 'ממוקד',
                                     'user_decision': {'answer_kind': 'delegated'}, 'crosses_red_line': True}]
    st['numbering_map'] = {'D06': 'סעיף 1'}
    st['produced_documents'] = [{'kind': 'work', 'path': 'C:/Users/x/deal/a.docx', 'sha256': ''}]
    st['deliverables'] = []
    codes = {e['code'] for e in sb.validate_state(st)}
    assert {'STATE_ROUND_STAGE', 'STATE_POLICY_GATE_TYPE', 'STATE_APPROVAL_FIELD', 'STATE_CLAUSE_REF',
            'STATE_RED_LINE_WITHOUT_USER_DECISION', 'STATE_NUMBERING_MAP_KEY', 'STATE_PRODUCED_PATH',
            'STATE_PRODUCED_SHA', 'STATE_UNKNOWN_KEY'} <= codes


def test_build_state_is_valid_and_has_no_approval_fields(tmp_path):
    out, exp, dec = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    deal = str(tmp_path)
    st = sb.build_state('deal_fixture_services2', 1, H.PROFILE, decisions=dec, deal_dir=deal,
                        produced=[('work', out)], clause_labels={'S04': 'אחריות, סעיף 1'})
    assert st['meta']['round_stage'] == 'opening' and st['numbering_map']['S04'] == 'אחריות, סעיף 1'
    sb.record_user_decision(st, 'Q2', 2, 'מקבלים', 'explicit', applies_to=['S04'])
    assert sb.validate_state(st) == []
    with pytest.raises(ValueError):
        sb.record_user_decision(st, 'Q10', 10, 'x', 'explicit')
    flat = json.dumps(st, ensure_ascii=False)
    assert '"approval"' not in flat and 'required_by' not in flat and 'authority_table' not in flat


def test_migrate_round3_state_to_49():
    old = {'meta': {'schema_version': '4.8', 'round_stage': 'awaiting_playbook_builder', 'round_number': 1},
           'playbook_status': {'playbook_resolution': {'authority_table': [{'source': 'x'}]}},
           'clauses': [{'clause_id': 'E-013', 'approval': {'status': 'proposed', 'required_by': 'מנכ"ל'}}],
           'deliberate_concessions': [{'topic': 't', 'clause_ref': 'E-013', 'reason': 'r', 'answer_kind': 'delegated',
                                       'approval': 'proposed'}],
           'deliverables': ['a.docx']}
    new, log = sm.migrate_state(old)
    assert new['meta']['round_stage'] == 'opening' and new['meta']['schema_version'] == '4.9'
    assert new['positions_source']['proposed_playbook']['build_state'] == 'awaiting_builder'
    assert new['clauses'][0]['user_decision']['answer_kind'] == 'delegated'
    assert new['deliberate_concessions'][0]['user_decision']['answer_kind'] == 'delegated'
    flat = json.dumps(new, ensure_ascii=False)
    assert 'required_by' not in flat and 'authority_table' not in flat and '"approval"' not in flat
    assert 'produced_documents' in new and 'deliverables' not in new


def test_question_batch_rules():
    items = [
        {'key': 'components', 'text': 'להריץ הערכת סיכונים לפני ההכרעות?', 'kind': 'components',
         'options': ['לפני', 'אחרי']},
        {'key': 'entity', 'text': 'מספר החברה בפרופיל אינו תקין.', 'situation': 3, 'blocks_sending': True,
         'options': ['למסור עכשיו', 'נשלים לפני החתימה']},
        {'key': 'law', 'text': 'דין קליפורניה וסנטה קלרה.', 'situation': 2, 'recommendation': 'לקבל עם סעד דחוף בישראל'},
        {'key': 'value', 'text': 'שווי ההתקשרות?', 'kind': 'deal_card_value'},
    ]
    b = qb.build_batch(items, no_consideration=True, asked_before={'components'})
    assert b['text'].splitlines()[0].startswith('1. **נדרשת הכרעתך (מצב 3, חוסם משלוח, תשובה של מילה אחת):**')
    assert 'נדרשת הכרעתך (מצב 2)' in b['text'] and 'מנכ' not in b['text'] and 'אישור' not in b['text']
    assert {x['key'] for x in b['skipped']} == {'components', 'value'}
    with pytest.raises(ValueError):
        qb.build_batch([{'key': 'x', 'text': 'y', 'situation': 10}])
    many = [{'key': 'k%d' % i, 'text': 'פער %d.' % i, 'situation': 1, 'recommendation': 'r'} for i in range(10)]
    assert len(qb.build_batch(many)['questions']) == 1


def test_review_table_has_no_approver_column():
    decs = [{'id': 'D22', 'summary': 'תקרה', 'source': 'playbook'},
            {'id': 'E-013', 'summary': 'דין קליפורניה', 'crosses_red_line': True, 'user_decision': {'answer_kind': 'explicit',
                                                                                                  'answer': 'מקבלים'}},
            {'id': 'R-19', 'summary': 'בקרת יצוא', 'situation': 8}]
    rows = rt.build_review_table(decs, {'D22': 'אחריות, סעיף 1'})
    assert [r['סעיף'] for r in rows] == ['אחריות, סעיף 1', 'E-013']
    assert rows[1]['תווית מקור'].startswith('חוצה קו אדום: הכרעתך')
    assert rt.check_review_table(rows, decs[:2]) == []
    assert 'מי מאשר' not in rt.render_markdown(rows)
    bad = [{'סעיף': 'x', 'מה הוכרע': 'y', 'תווית מקור': 'מהפלייבוק', 'מי מאשר': 'מנכ"ל'}]
    assert {'approval_chain_wording', 'extra_column:מי מאשר'} <= set(rt.check_review_table(bad))


def test_talking_points_and_signing_checklist(tmp_path):
    pts = tp.build_talking_points([{'id': 'D22', 'topic': 'תקרת האחריות', 'expected_friction': 'high',
                                    'prepared_fallback': 'הנסיגה שבפלייבוק', 'at_last_fallback': True}],
                                  concessions=[{'topic': 'נושא שוויתרנו עליו', 'if_raised': 'מה שמבקשים בתמורה'}])
    kinds = [p['kind'] for p in pts]
    assert kinds == ['friction', 'last_fallback', 'concession']
    assert 'הנסיגה שבפלייבוק' in pts[0]['text'] and 'נדרשת הכרעתך' in pts[1]['text']
    out, exp, _ = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    g = qa_gate.run_gate(out, S2, 'work', H.PROFILE, SIGNING, exp)
    chk = sc.build_signing_checklist(g['info'], SIGNING, H.PROFILE, register=_register())
    items = {c['item'] for c in chk['checks']}
    assert {'signing_blank', 'their_placeholder', 'registered_legal_name', 'pre_send_fix'} <= items
    assert chk['passed'] is False


# --------------------------------------------------------------------------------- (6) נתיב
def test_one_path_limit_230_in_gate_and_writer(tmp_path):
    # check_write_path מודד את הנתיב כפי שהוא על הדיסק (os.path.abspath),
    # ולכן הנתיב בבדיקה נבנה מוחלט בפלטפורמה שרצה. נתיב בסגנון Windows
    # אינו מוחלט על macOS ולינוקס, abspath מוסיף לפניו את תיקיית העבודה,
    # והאורך שנמדד אינו מה שהבדיקה התכוונה אליו
    assert paths.PATH_LIMIT == 230
    root = os.path.abspath(os.sep)
    pad = paths.PATH_LIMIT + 1 - len(root) - len('.docx')
    long = os.path.join(root, 'א' * pad + '.docx')
    assert len(os.path.abspath(long)) > paths.PATH_LIMIT
    with pytest.raises(paths.PathTooLong):
        paths.check_write_path(long)
    out, exp, _ = _apply(tmp_path, S2, 'services2_decisions_clean.json')
    g = qa_gate.run_gate(out, S2, 'work', H.PROFILE, {}, exp, work_path=long)
    assert 'PKG_PATH_LENGTH' in _codes(g) and '230' in _msgs(g, 'PKG_PATH_LENGTH')
    ok = os.path.join(root, 'א' * (pad - 11) + '.docx')
    assert len(os.path.abspath(ok)) <= paths.PATH_LIMIT
    assert 'PKG_PATH_LENGTH' not in _codes(qa_gate.run_gate(out, S2, 'work', H.PROFILE, {}, exp, work_path=ok))

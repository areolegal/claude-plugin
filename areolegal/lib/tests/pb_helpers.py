# -*- coding: utf-8 -*-
"""עזרי בדיקה לרכיב הפלייבוק: טעינת fixtures, חתימה מדומה על גיליון העמדות (כמו שהיועץ המשפטי עושה ב-Word),
ו-Word ו-Workspace קטנים של פלייבוק, כדי לבדוק את ההפקה מחדש אחרי הקליטה."""
import copy
import io
import json
import os
import re
import sys
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from common import rtl  # noqa: E402
from playbook.approval_sheet import HE  # noqa: E402

FIX = os.path.join(ROOT, 'tests', 'fixtures', 'playbook')
W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'


def load(name):
    with io.open(os.path.join(FIX, name), encoding='utf-8') as fh:
        return json.load(fh)


def good_model():
    return copy.deepcopy(load('mutual_nda_model_good.json'))


def round3_model():
    return copy.deepcopy(load('mutual_nda_round3_excerpt.json'))


def _check(doc, tags):
    for sdt in doc.element.body.iter(qn('w:sdt')):
        tag = sdt.find(qn('w:sdtPr') + '/' + qn('w:tag'))
        if tag is None or tag.get(qn('w:val')) not in tags:
            continue
        chk = sdt.find('.//{%s}checked' % W14)
        chk.set('{%s}val' % W14, '1')
        for t in sdt.iter(qn('w:t')):
            t.text = (t.text or '').replace('☐', '☒')


def _add(paragraph, text):
    """ריצה שאדם מקליד ב-Word במסמך עברי: עם rtl, כמו שהתבנית מגדירה."""
    r = paragraph.add_run(text)
    rtl.set_run_rtl(r)
    return r


def sign_sheet(src, dst, signers, choices=(), assumptions=(), legal_by=None, remove_areo_comments=False,
               final=True):
    """signers: {track: (name, role, date)}; choices / assumptions: תגיות תיבות לסמן;
    legal_by: {rule_id: name | (name, date)} לשורות "נבדק בידי" ו"תאריך הבדיקה"."""
    doc = Document(src)
    tags = set(choices) | set(assumptions) | {'chk:sign:%s' % t for t in signers}
    if final:
        tags.add('chk:final')
    _check(doc, tags)
    labels = {v: k for k, v in HE['track_label'].items()}
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            first = cells[0].paragraphs[0].text.strip() if cells[0].paragraphs else ''
            if len(cells) == 5 and labels.get(first) in signers:
                name, role, date = signers[labels[first]]
                for c, v in zip(cells[2:5], (name, role, date)):
                    _add(c.paragraphs[0], v)
            elif len(cells) == 4 and legal_by and first in legal_by:
                who = legal_by[first]
                name, date = (who, '14.10.2026') if isinstance(who, str) else who
                for c in cells[1:]:
                    for p in c.paragraphs:
                        if p.text.startswith(HE['legal_line'].strip()):
                            _add(p, name)
                        elif p.text.startswith(HE['legal_date'].strip()) and date:
                            _add(p, date)
    doc.save(dst)
    if remove_areo_comments:
        _strip_comments(dst)
    return dst


def _strip_comments(path):
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    doc = items['word/document.xml'].decode('utf-8')
    doc = re.sub(r'<w:commentRangeStart [^>]*/>|<w:commentRangeEnd [^>]*/>', '', doc)
    doc = re.sub(r'<w:r>(?:(?!</w:r>).)*?<w:commentReference [^>]*/>(?:(?!</w:r>).)*?</w:r>', '', doc, flags=re.S)
    items['word/document.xml'] = doc.encode('utf-8')
    if 'word/comments.xml' in items:
        c = items['word/comments.xml'].decode('utf-8')
        c = re.sub(r'<w:comment [^>]*w:author="Areo".*?</w:comment>', '', c, flags=re.S)
        items['word/comments.xml'] = c.encode('utf-8')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, data in items.items():
            z.writestr(n, data)


SIGNER = ('דנה לוי', 'יועצת משפטית ראשית', '14.10.2026')
ALL_SIGNERS = {'green': SIGNER, 'yellow': SIGNER, 'red': SIGNER}


# ------------------------------------------------------------------ Word ו-Workspace של פלייבוק
def playbook_docx(path, model, label, tagged=True, cell_lines=None):
    """Word קטן של פלייבוק. tagged: לפי החוזה של surfaces_render (בקרי תוכן). אחרת: כמו ה-Word של סבב 3,
    טקסט חופשי בכותרת העליונה, בתמונת המצב וליד כל מסלול. cell_lines: מספר שורות "מעמד העמדה" (ברירת מחדל:
    שלוש לכל כלל)."""
    from playbook import status as S
    from playbook import surfaces_render as R
    doc = Document()
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    if tagged:
        R.add_status_control(hp, R.STATUS_TAG, label)
    else:
        _add(hp, label)
    rtl.add_paragraph(doc, 'פלייבוק: %s' % model.get('family_label', ''), bold=True, size=14)
    p = doc.add_paragraph()
    if tagged:
        R.add_status_control(p, R.SUMMARY_TAG, R.SUMMARY_PREFIX + label)
    else:
        _add(p, 'מעמד העמדות: גיליון העמדות נחתם ונקלט בסימולציה, ולכן אף עמדה אינה נספרת כמאושרת. 0 מתוך 12 '
                'עמדות מאושרות.')
    keys = list(S.iter_cell_keys(model))
    if cell_lines is not None:
        keys = keys[:cell_lines]
    for k in keys:
        rtl.add_paragraph(doc, k, size=10)
        p = doc.add_paragraph()
        if tagged:
            R.add_status_control(p, R.CELL_TAG + k, R.CELL_PREFIX + 'טעונה אישור')
        else:
            _add(p, R.CELL_PREFIX + 'לא מאושרת (החתימה נקלטה בסימולציה)')
    disc = (model.get('market_research') or {}).get('disclosure', {}).get('fixed_sentence')
    if disc:
        rtl.add_paragraph(doc, disc, size=10)
    rtl.enforce_document(doc, author='EXAMPLECO', title='פלייבוק')
    doc.save(path)
    return path


def workspace_html(path, model, label, disclaimer=None):
    """Workspace קטן: תבנית עם נתונים מוזרקים כ-const D = {...};, כמו ב-C2."""
    meta = {'title': 'פלייבוק: %s' % model.get('family_label', ''), 'status': label, 'presentation_label': label,
            'label_now': label, 'approved_count': 0, 'positions_count': 12,
            'evidence_disclaimer': disclaimer if disclaimer is not None else label + '. ' + (
                (model.get('market_research') or {}).get('disclosure', {}).get('fixed_sentence') or '')}
    data = {'meta': meta, 'rules': [{'id': r['rule_id']} for r in model.get('rules') or []]}
    html = ('<!doctype html><html lang="he" dir="rtl"><head><meta charset="utf-8"><title>Areo Workspace</title>'
            '</head><body><div id="app"></div><script>\nconst D = %s;\nrender(D);\n</script></body></html>'
            % json.dumps(data, ensure_ascii=False))
    with io.open(path, 'w', encoding='utf-8') as fh:
        fh.write(html)
    return path
